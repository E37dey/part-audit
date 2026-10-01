"""Excel report (openpyxl).

Sheets: Summary, Duplicate_Groups, Near_Duplicates, Name_Inconsistencies,
PartNumber_Mapping, ERP_vs_CAD, Errors_Skipped, Metrics (demo only).
"""
from __future__ import annotations

import datetime as _dt
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import __version__
from .names import pn_key

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
RED_FILL = PatternFill("solid", fgColor="F8D7DA")
AMBER_FILL = PatternFill("solid", fgColor="FFE8B3")
GREEN_FILL = PatternFill("solid", fgColor="D4EDDA")
WRAP = Alignment(wrap_text=True, vertical="top")

CATEGORY_LABEL = {
    "identical": "Identical (merge candidate)",
    "identical_attr_conflict": "Identical geometry, ERP attributes differ (do not merge)",
    "near_identical": "Near-identical (alert, do not merge)",
    "conflict": "Conflict: same part number, different geometry",
    "conflict_unverified": "Same part number, geometry not compared",
    "name_similar": "Similar names only",
    "mirror": "Mirror image (left/right-hand part) - not a duplicate",
}


def _pct(x):
    if x is None or (isinstance(x, float) and (math.isinf(x) or math.isnan(x))):
        return None
    return float(x)


def _write_table(ws, rows: list[dict], columns: list[tuple[str, str, int]], start_row: int = 1,
                 pct_cols: tuple[str, ...] = (), wrap_cols: tuple[str, ...] = ()) -> int:
    """Write rows as a styled table; columns = [(key, header, width)]. Returns last row."""
    for ci, (_, header, width) in enumerate(columns, start=1):
        c = ws.cell(row=start_row, column=ci, value=header)
        c.fill, c.font, c.alignment = HEADER_FILL, HEADER_FONT, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(ci)].width = width
    r = start_row
    for row in rows:
        r += 1
        for ci, (key, _, _) in enumerate(columns, start=1):
            v = row.get(key)
            if isinstance(v, float) and (math.isinf(v) or math.isnan(v)):
                v = None
            c = ws.cell(row=r, column=ci, value=v)
            if key in pct_cols and isinstance(v, (int, float)):
                c.number_format = "0.0000%"
            if key in wrap_cols:
                c.alignment = WRAP
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(columns))}{r}"
    else:
        ws.cell(row=start_row + 1, column=1, value="(none)").font = Font(italic=True, color="666666")
    return r


def _item_name(it) -> str:
    return it.rec.stem


def _short(h: str) -> str:
    return h[:12] if h else ""


def build_rows(res) -> dict[str, list[dict]]:
    """Plain row dicts for every sheet (also used by tests and CSV export)."""
    items = res.items
    pair_by_key = {(p.a, p.b): p for p in res.pairs}
    out: dict[str, list[dict]] = {}

    # Duplicate groups
    map_canon = {}
    for m in res.mapping:
        map_canon[m["group_id"]] = m["proposed_pn"]
    dg = []
    for gi, g in enumerate(res.groups, start=1):
        gid = f"G{gi:04d}"
        gpairs = [p for (a, b), p in pair_by_key.items() if a in g and b in g and p.category == "identical"]
        conf = min((p.confidence for p in gpairs), default=None)
        kinds = sorted({("byte-identical" if p.l1 == "raw" else "normalized STEP" if p.l1 == "norm"
                         else "unit-scaled geometry" if p.geo_relation == "unit_scaled" else "geometry")
                        for p in gpairs})
        for fid in g:
            it = items[fid]
            best = max((p for p in gpairs if fid in (p.a, p.b)), key=lambda p: p.confidence, default=None)
            fp = it.fp or {}
            dg.append({
                "group_id": gid, "group_size": len(g), "file_id": fid, "folder": it.rec.folder,
                "file_name": it.rec.filename, "part_number": it.pn, "erp_description": (it.erp or {}).get("description", ""),
                "erp_material": (it.erp or {}).get("material", ""),
                "step_product": it.rec.products[0][1] if it.rec.products else "",
                "declared_unit": it.rec.length_unit, "sha256": _short(it.rec.sha256),
                "volume_mm3": fp.get("volume"), "area_mm2": fp.get("area"), "faces": fp.get("n_faces"),
                "match_type": ", ".join(kinds), "group_confidence": conf,
                "proposed_pn": map_canon.get(gid, it.pn), "evidence": best.reason if best else "",
            })
    out["Duplicate_Groups"] = dg

    # Near duplicates and geometry conflicts
    nd = []
    for p in res.pairs:
        if p.category in ("near_identical", "conflict", "identical_attr_conflict", "mirror"):
            g = p.geo
            shp = (g.shape if g else None) or {}
            ia, ib = items[p.a], items[p.b]
            label = CATEGORY_LABEL.get(p.category, p.category)
            if p.category == "conflict" and p.geo_relation == "near":
                label = "Conflict: same part number, slightly different geometry (revision?)"
            nd.append({
                "category": label, "geo_relation": p.geo_relation,
                "file_a": p.a, "file_b": p.b, "pn_a": ia.pn, "pn_b": ib.pn, "confidence": p.confidence,
                "d_volume": _pct(g.d_volume) if g else None, "d_area": _pct(g.d_area) if g else None,
                "d_inertia": _pct(g.d_inertia) if g else None, "d_bbox": _pct(g.d_bbox) if g else None,
                "faces_a": (ia.fp or {}).get("n_faces"), "faces_b": (ib.fp or {}).get("n_faces"),
                "surface_p95": shp.get("p95_rel"), "reason": p.reason,
            })
    nd.sort(key=lambda r: (r["category"], -(r["confidence"] or 0)))
    out["Near_Duplicates"] = nd

    # Name inconsistencies
    ni = []
    for p in res.pairs:
        ia, ib = items[p.a], items[p.b]
        na = ia.sources[0] if ia.sources else None
        nb = ib.sources[0] if ib.sources else None
        base = {"file_a": p.a, "file_b": p.b, "name_a": _item_name(ia), "name_b": _item_name(ib),
                "normalized_a": na.display if na else "", "normalized_b": nb.display if nb else "",
                "pn_a": ia.pn, "pn_b": ib.pn, "name_score": p.name.score if p.name else None,
                "geometry": p.geo_relation, "confidence": p.confidence}
        if p.category == "identical" and (ia.pn_key != ib.pn_key or (na and nb and na.canon != nb.canon)):
            issue = ("Same part, different part numbers" if ia.pn_key and ib.pn_key and ia.pn_key != ib.pn_key
                     else "Same part, part number missing in one name" if bool(ia.pn_key) != bool(ib.pn_key)
                     else "Same part, different descriptions")
            ni.append({**base, "issue": issue, "suggestion": "use one part number and one description; see PartNumber_Mapping"})
        elif p.category == "conflict":
            ni.append({**base, "issue": "Same part number, different geometry",
                       "suggestion": "assign a new part number or revision to one of the files; " + p.reason})
        elif p.category == "conflict_unverified":
            ni.append({**base, "issue": "Same part number, geometry not compared", "suggestion": p.reason})
        elif p.category == "name_similar":
            ni.append({**base, "issue": ("Similar names, different parts" if p.geo_relation == "different"
                                         else "Similar names, geometry not compared"),
                       "suggestion": ("make descriptions distinguishable" if p.geo_relation == "different"
                                      else "check manually whether these are the same part")})
    for it in items.values():  # file-name PN spelled differently from the ERP PN
        if it.erp and it.sources and it.sources[0].pns:
            raw = it.sources[0].pns[0]
            if raw != it.erp["part_number"] and pn_key(raw) == pn_key(it.erp["part_number"]):
                ni.append({"issue": "PN spelled differently from ERP", "file_a": it.file_id, "file_b": "",
                           "name_a": _item_name(it), "name_b": f"ERP: {it.erp['part_number']}",
                           "normalized_a": it.sources[0].display, "normalized_b": pn_key(it.erp["part_number"]),
                           "pn_a": raw, "pn_b": it.erp["part_number"], "name_score": 1.0, "geometry": "",
                           "confidence": 1.0, "suggestion": f"rename file to use '{it.erp['part_number']}'"})
    order = {"Same part number, different geometry": 0, "Same part, different part numbers": 1,
             "Same part, part number missing in one name": 2, "Same part, different descriptions": 3,
             "Same part number, geometry not compared": 4, "Similar names, geometry not compared": 5,
             "Similar names, different parts": 6, "PN spelled differently from ERP": 7}
    ni.sort(key=lambda r: (order.get(r["issue"], 9), -(r.get("name_score") or 0)))
    out["Name_Inconsistencies"] = ni

    # Mapping
    llm = (res.llm_notes or {}).get("mapping", {})
    out["PartNumber_Mapping"] = [{**m, "approved": None, "reviewer_notes": None, "llm_note": llm.get(i)}
                                 for i, m in enumerate(res.mapping)]

    # ERP vs CAD
    ev = []
    files_by_erp: dict[str, list[str]] = {}
    for it in items.values():
        if it.erp:
            files_by_erp.setdefault(it.erp["part_number"], []).append(it.file_id)
    group_of = {fid: f"G{gi:04d}" for gi, g in enumerate(res.groups, start=1) for fid in g}
    mapped_old = {m["old_pn"]: m["proposed_pn"] for m in res.mapping}
    for row in res.erp_rows:
        pn = row["part_number"]
        fids = sorted(files_by_erp.get(pn, []))
        if not fids:
            status, detail = "No CAD file found", "no STEP file name/product matches this part number"
        elif pn in mapped_old:
            status, detail = "Duplicate of another ERP item", f"same geometry as {mapped_old[pn]} (see PartNumber_Mapping)"
        elif len(fids) > 1 and len({group_of.get(f, f) for f in fids}) > 1:
            status, detail = "Several different CAD files", "files with this PN are not identical (see Name_Inconsistencies)"
        else:
            status, detail = "OK", ""
        ev.append({"source": "ERP", "part_number": pn, "erp_row": row.get("erp_row"),
                   "description": row.get("description", ""), "unit": row.get("unit", ""),
                   "material": row.get("material", ""), "supplier": row.get("supplier", ""),
                   "cad_files": len(fids), "cad_file_ids": "; ".join(fids[:10]) + (" ..." if len(fids) > 10 else ""),
                   "status": status, "detail": detail})
    for it in sorted(items.values(), key=lambda x: x.file_id):
        if not it.erp:
            ev.append({"source": "CAD", "part_number": it.pn, "erp_row": None, "description": _item_name(it),
                       "unit": "", "material": "", "supplier": "", "cad_files": 1, "cad_file_ids": it.file_id,
                       "status": "CAD file not in ERP" if res.erp_rows else "no ERP file given",
                       "detail": "no ERP part number matched the file name or STEP product"})
    for d in res.erp_duplicates:
        ev.append({"source": "ERP", "part_number": d["pn_a"], "erp_row": d["erp_row_a"], "description": d["desc_a"],
                   "unit": "", "material": "", "supplier": "", "cad_files": None, "cad_file_ids": "",
                   "status": "Possible duplicate ERP item (by description)",
                   "detail": f"row {d['erp_row_b']} {d['pn_b']} '{d['desc_b']}' - {d['detail']}"})
    out["ERP_vs_CAD"] = ev

    out["Errors_Skipped"] = [{"file_id": s.file_id, "stage": s.stage, "reason": s.reason, "detail": s.detail,
                              "size_kb": round(s.size / 1024, 1)} for s in res.skipped]
    return out


def write_report(res, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = build_rows(res)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _summary(ws, res, rows)

    ws = wb.create_sheet("Duplicate_Groups")
    _write_table(ws, rows["Duplicate_Groups"], [
        ("group_id", "Group", 9), ("group_size", "Files", 6), ("file_id", "File", 48), ("part_number", "Part number", 16),
        ("proposed_pn", "Proposed PN", 16), ("erp_description", "ERP description", 28), ("erp_material", "ERP material", 14),
        ("step_product", "STEP product name", 22), ("declared_unit", "Unit in file", 8), ("match_type", "Match type", 20),
        ("group_confidence", "Confidence", 10), ("volume_mm3", "Volume mm³", 13), ("area_mm2", "Area mm²", 12),
        ("faces", "Faces", 7), ("sha256", "SHA-256", 14), ("evidence", "Evidence", 70)], wrap_cols=("evidence",))
    _band_groups(ws, rows["Duplicate_Groups"])

    ws = wb.create_sheet("Near_Duplicates")
    _write_table(ws, rows["Near_Duplicates"], [
        ("category", "Category", 30), ("file_a", "File A", 42), ("file_b", "File B", 42), ("pn_a", "PN A", 14),
        ("pn_b", "PN B", 14), ("confidence", "Confidence", 10), ("d_volume", "Δ volume", 10), ("d_area", "Δ area", 10),
        ("d_inertia", "Δ inertia", 10), ("d_bbox", "Δ bbox", 10), ("faces_a", "Faces A", 8), ("faces_b", "Faces B", 8),
        ("surface_p95", "Surface dist p95", 11), ("reason", "Reason", 80)],
        pct_cols=("d_volume", "d_area", "d_inertia", "d_bbox", "surface_p95"), wrap_cols=("reason",))

    ws = wb.create_sheet("Name_Inconsistencies")
    _write_table(ws, rows["Name_Inconsistencies"], [
        ("issue", "Issue", 32), ("file_a", "File A", 40), ("file_b", "File B", 40), ("name_a", "Name A", 28),
        ("name_b", "Name B", 28), ("normalized_a", "Normalized A", 26), ("normalized_b", "Normalized B", 26),
        ("pn_a", "PN A", 14), ("pn_b", "PN B", 14), ("name_score", "Name score", 9), ("geometry", "Geometry", 11),
        ("confidence", "Confidence", 10), ("suggestion", "Suggestion", 60)], wrap_cols=("suggestion",))
    _highlight_issue(ws, rows["Name_Inconsistencies"])

    ws = wb.create_sheet("PartNumber_Mapping")
    ws.cell(row=1, column=1, value="Proposed part-number mapping - nothing is changed in the ERP. "
                                   "Fill 'Approved (Y/N)' and 'Reviewer notes'.").font = BOLD
    map_cols = [
        ("group_id", "Group", 9), ("old_pn", "Old PN", 16), ("proposed_pn", "Proposed PN", 16),
        ("reason", "Reason", 70), ("confidence", "Confidence", 10), ("status", "Status", 16),
        ("approved", "Approved (Y/N)", 13), ("reviewer_notes", "Reviewer notes", 36),
        ("files_old", "Files (old PN)", 44), ("files_proposed", "Files (proposed PN)", 44)]
    if any(r.get("llm_note") for r in rows["PartNumber_Mapping"]):
        map_cols.append(("llm_note", "LLM note (advisory, unverified)", 60))
    last = _write_table(ws, rows["PartNumber_Mapping"], map_cols, start_row=3,
                        wrap_cols=("reason", "files_old", "files_proposed", "llm_note"))
    if rows["PartNumber_Mapping"]:
        dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"G4:G{last}")
        for r in range(4, last + 1):
            ws.cell(row=r, column=7).fill = INPUT_FILL
            ws.cell(row=r, column=8).fill = INPUT_FILL

    ws = wb.create_sheet("ERP_vs_CAD")
    _write_table(ws, rows["ERP_vs_CAD"], [
        ("source", "Source", 7), ("part_number", "Part number", 16), ("erp_row", "ERP row", 8),
        ("description", "Description", 32), ("unit", "Unit", 7), ("material", "Material", 16), ("supplier", "Supplier", 16),
        ("cad_files", "CAD files", 9), ("cad_file_ids", "CAD file(s)", 48), ("status", "Status", 30),
        ("detail", "Detail", 60)], wrap_cols=("detail",))

    ws = wb.create_sheet("Errors_Skipped")
    _write_table(ws, rows["Errors_Skipped"], [
        ("file_id", "File", 60), ("stage", "Stage", 10), ("reason", "Reason", 30), ("detail", "Detail", 70),
        ("size_kb", "Size KB", 10)], wrap_cols=("detail",))

    if res.metrics:
        ws = wb.create_sheet("Metrics")
        _metrics_sheet(ws, res.metrics)
    for w in wb.worksheets:
        w.sheet_view.zoomScale = 90
        _print_setup(w, None if w.title in ("Summary", "Metrics") else (3 if w.title == "PartNumber_Mapping" else 1))
    wb.save(path)
    return path


SECTION_FONT = Font(bold=True, size=12, color="1F3A5F")
GREY_FONT = Font(color="666666", size=9)


def _summary(ws, res, rows) -> None:
    cats: dict[str, int] = {}
    for p in res.pairs:
        cats[p.category] = cats.get(p.category, 0) + 1
    n_files = len(res.items)
    n_in_groups = sum(len(g) for g in res.groups)
    geo_err = sum(1 for s in res.skipped if s.stage == "geometry")
    ws["A1"] = "Part Data Audit - summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ("Read-only audit: nothing was written to any ERP or CAD system. Every proposed part-number "
                "change is pending human approval in the PartNumber_Mapping sheet.")
    ws.merge_cells("A2:C2")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30

    findings = [
        ("STEP files analysed", n_files, ""),
        ("Files skipped or failed", len(res.skipped), f"{geo_err} geometry errors/timeouts - see Errors_Skipped"),
        ("ERP rows", len(res.erp_rows), ""),
        ("Duplicate groups (identical parts)", len(res.groups), "Duplicate_Groups"),
        ("Files in duplicate groups", n_in_groups, f"{n_in_groups - len(res.groups)} could be retired after review"),
        ("Proposed part-number mappings", len(res.mapping), "PartNumber_Mapping - pending approval"),
        ("Near-identical pairs (alert, never merged)", cats.get("near_identical", 0), "Near_Duplicates"),
        ("Part-number conflicts (same PN, different geometry)", cats.get("conflict", 0), "Name_Inconsistencies"),
        ("Identical geometry, but materials differ", cats.get("identical_attr_conflict", 0), "Near_Duplicates"),
        ("Mirror-image pairs (left/right, not duplicates)", cats.get("mirror", 0), "Near_Duplicates"),
        ("Same PN in files whose geometry could not be compared", cats.get("conflict_unverified", 0), ""),
        ("Similar names, different or unknown geometry", cats.get("name_similar", 0), "Name_Inconsistencies"),
        ("Name inconsistency rows", len(rows["Name_Inconsistencies"]), "Name_Inconsistencies"),
        ("ERP items without a CAD file", sum(1 for r in rows["ERP_vs_CAD"] if r["status"] == "No CAD file found"), "ERP_vs_CAD"),
        ("CAD files not in the ERP", sum(1 for r in rows["ERP_vs_CAD"] if r["status"] == "CAD file not in ERP"), "ERP_vs_CAD"),
        ("Possible duplicate ERP items (same description)", len(res.erp_duplicates), "ERP_vs_CAD"),
    ]
    how = [
        "1. Duplicate_Groups: each group is one physical part saved several times; check the evidence column.",
        "2. PartNumber_Mapping: for groups with several part numbers, approve (Y) or reject (N) the proposed PN.",
        "3. Near_Duplicates: small design changes, other-material copies and mirror images - review, never merge blindly.",
        "4. Name_Inconsistencies: same part under different names/PNs, and same PN used for different geometry.",
        "5. ERP_vs_CAD: ERP items without CAD and CAD files without an ERP item.",
    ]
    run = [
        ("Scanned folder", res.config.root), ("ERP file", res.config.erp_csv or "(none)"),
        ("Run date", _dt.datetime.now().strftime("%Y-%m-%d %H:%M")), ("Tool version", __version__),
        ("Geometry layer (L3)", "on" if res.config.use_geometry else "off (names and hashes only)"),
        ("Shape check (L3b)", "on" if res.config.use_geometry and res.config.shape_check else "off"),
        ("LLM layer (L4)", ("on: " + res.config.llm_endpoint) if res.config.llm_endpoint else "off (default, no network use)"),
        ("Worker processes", res.config.workers), ("Per-file timeout (s)", res.config.timeout_s),
    ]
    r = 4
    ws.cell(row=r, column=1, value="Findings").font = SECTION_FONT
    r += 1
    for k, v, where in findings:
        ws.cell(row=r, column=1, value=k).font = BOLD
        ws.cell(row=r, column=2, value=v)
        ws.cell(row=r, column=3, value=where).font = GREY_FONT
        ws.cell(row=r, column=3).alignment = Alignment(indent=1)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="How to use this report").font = SECTION_FONT
    r += 1
    for line in how:
        ws.cell(row=r, column=1, value=line)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Run").font = SECTION_FONT
    r += 1
    for k, v in run:
        ws.cell(row=r, column=1, value=k).font = BOLD
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Timings (seconds) and thresholds").font = SECTION_FONT
    r += 1
    for k, v in sorted(res.timings.items()):
        ws.cell(row=r, column=1, value=f"time: {k}").font = GREY_FONT
        ws.cell(row=r, column=2, value=round(v, 2) if isinstance(v, float) else v).font = GREY_FONT
        r += 1
    for k, v in res.config.thresholds.as_dict().items():
        ws.cell(row=r, column=1, value=f"threshold: {k}").font = GREY_FONT
        ws.cell(row=r, column=2, value=v).font = GREY_FONT
        r += 1
    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 60
    for row in ws.iter_rows(min_row=5, max_row=4 + len(findings), min_col=2, max_col=2):
        for c in row:
            c.alignment = Alignment(horizontal="right")


def _print_setup(ws, header_row: int | None) -> None:
    """Landscape A4, fit to one page wide, header row repeated - so the report prints usefully."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    if header_row:
        ws.print_title_rows = f"{header_row}:{header_row}"


def _band_groups(ws, rows) -> None:
    fills = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="EEF3F8")]
    gid, k = None, 0
    for i, row in enumerate(rows, start=2):
        if row["group_id"] != gid:
            gid, k = row["group_id"], k + 1
        for c in range(1, 8):
            ws.cell(row=i, column=c).fill = fills[k % 2]


def _highlight_issue(ws, rows) -> None:
    for i, row in enumerate(rows, start=2):
        fill = RED_FILL if row["issue"].startswith("Same part number, different") else \
            AMBER_FILL if row["issue"].startswith("Same part,") else None
        if fill:
            ws.cell(row=i, column=1).fill = fill


def _metrics_sheet(ws, metrics: dict) -> None:
    ws["A1"] = "Metrics against ground_truth.csv (DEMO ONLY - synthetic data, accuracy is an upper bound)"
    ws["A1"].font = TITLE_FONT
    r = 3
    table = metrics.get("per_type", [])
    cols = [("type", "Type", 30), ("n_pairs", "GT pairs", 9), ("tp", "TP", 6), ("fn", "FN", 6), ("fp", "FP", 6),
            ("precision", "Precision", 10), ("recall", "Recall", 10), ("f1", "F1", 8), ("target", "Target", 30),
            ("met", "Met?", 7), ("definition", "How it is counted", 70)]
    r = _write_table(ws, table, cols, start_row=r, wrap_cols=("definition",))
    for i in range(4, r + 1):
        for c in (6, 7, 8):
            ws.cell(row=i, column=c).number_format = "0.0%"
        met = ws.cell(row=i, column=10).value
        if met in ("yes", "no"):
            ws.cell(row=i, column=10).fill = GREEN_FILL if met == "yes" else RED_FILL
    r += 2
    for k, v in metrics.get("summary", {}).items():
        ws.cell(row=r, column=1, value=k).font = BOLD
        ws.cell(row=r, column=2, value=v)
        r += 1
    ws.freeze_panes = None
