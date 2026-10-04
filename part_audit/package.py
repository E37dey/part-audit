"""Change package (the engineering change order of a cleanup): what has to change in each system to carry out
the approved duplicate decisions. PartAudit only prepares the list - it never writes to a PDM, PLM or ERP.

For every approved row "old part number -> the part number that stays":
* every BOM line that uses the old part number (in the PDM / PLM, and in the ERP's BOM) is replaced, in the system
  it is in - an assembly that already has the part number that stays gets the quantities added;
* every record with the old part number is marked as replaced (obsolete, superseded by the one that stays);
* the old ERP item's stock is moved to the item that stays, and the old item is closed for new purchases.

The data is prepared per group and per possible target (`package_data`), so the online copy can put a package
together from the decisions made in the browser with the same rules (`combine`).
"""
from __future__ import annotations

from collections import defaultdict

from .decisions import chosen_target, effective_mapping, main_currency, row_decision
from .names import pn_key

ERP_KINDS = {"priority", "sap_s4", "sap_b1", "erp_file"}
KIND_ORDER = {"replace_line": 0, "supersede": 1, "move_stock": 2}


def _erp_name(res: dict) -> str:
    for s in res.get("sources", []):
        if s.get("kind") in ERP_KINDS:
            return s["name"]
    return "ERP"


def _key(pn: str | None) -> str:
    return pn_key(pn or "")


def package_data(res: dict) -> dict:
    """Everything a package can be made of: per group and per member that could be its target, the actions of each
    old part number and the BOM lines it replaces; and the BOMs of the assemblies those lines are in."""
    idx = {p["uid"]: p for p in res["parts"]}
    erp_name = _erp_name(res)
    erp_bom: dict[str, list[dict]] = defaultdict(list)
    for b in res.get("erp_bom", []):
        erp_bom[_key(b["parent"])].append(b)
    erp_desc = {}
    for p in res["parts"]:
        e = p.get("erp") or {}
        if e.get("pn"):
            erp_desc.setdefault(_key(e["pn"]), e.get("description") or p["name"])
    assemblies: dict[str, dict] = {}

    def pdm_assembly(uid: str) -> str:
        key = f"pdm:{uid}"
        if key not in assemblies and uid in idx:
            a = idx[uid]
            assemblies[key] = {"key": key, "system": a["source"], "system_kind": a["source_kind"], "kind": "pdm", "pn": a["pn"],
                               "name": a["name"], "uid": uid, "state": a.get("state"),
                               "lines": [{"pn": idx[c["child"]]["pn"], "name": idx[c["child"]]["name"], "qty": c.get("qty") or 0}
                                         for c in res.get("bom", {}).get(uid, []) if c["child"] in idx]}
        return key

    def erp_assembly(parent_pn: str, name: str) -> str:
        key = f"erp:{_key(parent_pn)}"
        if key not in assemblies:
            assemblies[key] = {"key": key, "system": erp_name, "system_kind": "ERP", "kind": "erp", "pn": parent_pn, "name": name, "uid": None,
                               "state": None, "lines": [{"pn": b["child"], "name": erp_desc.get(_key(b["child"]), ""),
                                                         "qty": b.get("qty") or 0} for b in erp_bom.get(_key(parent_pn), [])]}
        return key

    groups = {}
    for g in res.get("groups", []):
        targets = {}
        members = [idx[u] for u in g["members"] if u in idx]
        for t in members:
            if not t["pn"]:
                continue
            t_erp = (t.get("erp") or {}).get("pn")
            keep, keep_erp = _key(t["pn"]), _key(t_erp or t["pn"])
            rows: dict[str, dict] = {}
            for r in members:
                if not r["pn"] or _key(r["pn"]) == keep:
                    continue
                row = rows.setdefault(_key(r["pn"]), {"old_pn": r["pn"], "actions": [], "replace": []})
                row["actions"].append({"kind": "supersede", "system": r["source"], "system_kind": r["source_kind"],
                                       "pn": r["pn"], "name": r["name"], "uid": r["uid"], "state": r.get("state"),
                                       "new_pn": t["pn"]})
            for e in g.get("erp", []):
                if _key(e["pn"]) == keep_erp:
                    continue
                old = next((rows[k] for k in rows if k == _key(e["pn"])), None)
                if old is None:
                    # an ERP item of the group whose own record carries another number: file it under that record
                    owner = next((m for m in members if _key((m.get("erp") or {}).get("pn")) == _key(e["pn"])), None)
                    if owner is None or _key(owner["pn"]) == keep:
                        continue
                    old = rows.setdefault(_key(owner["pn"]), {"old_pn": owner["pn"], "actions": [], "replace": []})
                old["actions"].append({"kind": "move_stock", "system": erp_name, "system_kind": "ERP", "pn": e["pn"],
                                       "name": e.get("description") or "", "to_pn": t_erp or t["pn"], "stock": e.get("stock"),
                                       "uom": e.get("uom") or "", "value": e.get("value"), "currency": e.get("currency"),
                                       "unit_cost": e.get("unit_cost")})
            for w in g.get("where_used", []):
                if w.get("erp_pn"):
                    if _key(w["erp_pn"]) == keep_erp:
                        continue
                    owner = next((m for m in members if _key((m.get("erp") or {}).get("pn") or m["pn"]) == _key(w["erp_pn"])), None)
                    if owner is None or _key(owner["pn"]) == keep:
                        continue
                    a = erp_assembly(w["pn"], w.get("name") or "")
                    row = rows.setdefault(_key(owner["pn"]), {"old_pn": owner["pn"], "actions": [], "replace": []})
                    row["actions"].append({"kind": "replace_line", "system": erp_name, "system_kind": "ERP", "assembly": a,
                                           "assembly_pn": w["pn"], "assembly_name": w.get("name") or "", "old_pn": w["erp_pn"],
                                           "new_pn": t_erp or t["pn"], "qty": w.get("qty")})
                    row["replace"].append({"assembly": a, "old": w["erp_pn"], "new": t_erp or t["pn"], "new_name": t["name"]})
                else:
                    if not w.get("via") or _key(w["via"]) == keep or not w.get("uid"):
                        continue
                    a = pdm_assembly(w["uid"])
                    row = rows.setdefault(_key(w["via"]), {"old_pn": w["via"], "actions": [], "replace": []})
                    row["actions"].append({"kind": "replace_line", "system": w["source"], "system_kind": idx[w["uid"]]["source_kind"]
                                           if w["uid"] in idx else "", "assembly": a, "assembly_pn": w["pn"],
                                           "assembly_name": w.get("name") or "", "assembly_state": w.get("state"),
                                           "old_pn": w["via"], "new_pn": t["pn"], "qty": w.get("qty")})
                    row["replace"].append({"assembly": a, "old": w["via"], "new": t["pn"], "new_name": t["name"]})
            targets[t["uid"]] = {"target_pn": t["pn"], "target_erp_pn": t_erp, "target_name": t["name"],
                                 "target_source": t["source"], "rows": list(rows.values())}
        groups[g["id"]] = {"name": g["name"], "spec": g.get("spec"), "currency": g.get("currency"), "targets": targets}
    return {"erp_system": erp_name, "groups": groups, "assemblies": assemblies,
            "currency": res.get("summary", {}).get("currency") or "ILS"}


def approved(res: dict, decisions: dict) -> list[tuple[str, str, list[str]]]:
    """(group, target member, old part numbers) for every approved mapping row - a group approval approves all its
    rows; a row can also be approved, or excepted, on its own (decisions.row_decision)."""
    idx = {p["uid"]: p for p in res["parts"]}
    groups = {g["id"]: g for g in res.get("groups", [])}
    out: dict[str, tuple[str, list[str]]] = {}
    for r in effective_mapping(res, decisions):
        if row_decision(r, decisions)["decision"] != "approve" or r["group_id"] not in groups:
            continue
        g = groups[r["group_id"]]
        target = chosen_target(g, decisions.get(f"group:{g['id']}"), idx)
        if _key(idx[target]["pn"]) != _key(r["new_pn"]):        # a row approved for a target that is no longer chosen
            continue
        out.setdefault(g["id"], (target, []))[1].append(r["old_pn"])
    return [(gid, t, olds) for gid, (t, olds) in out.items()]


def combine(data: dict, chosen: list[tuple[str, str, list[str]]]) -> dict:
    """The package for the chosen (group, target, old part numbers): the actions per system and every affected
    assembly's BOM before and after."""
    actions, groups_out, repl = [], [], defaultdict(list)
    for gid, target, olds in chosen:
        g = data["groups"].get(gid)
        tg = g and g["targets"].get(target)
        if not tg:
            continue
        keys = {_key(o) for o in olds}
        rows = [r for r in tg["rows"] if _key(r["old_pn"]) in keys]
        groups_out.append({"id": gid, "name": g["name"], "spec": g.get("spec"), "target_pn": tg["target_pn"],
                           "target_name": tg["target_name"], "old_pns": [r["old_pn"] for r in rows]})
        for r in rows:
            for a in r["actions"]:
                actions.append({**a, "group_id": gid})
            for x in r["replace"]:
                repl[x["assembly"]].append(x)
    assemblies = []
    for key, rs in repl.items():
        a = data["assemblies"][key]
        assemblies.append({**{k: a[k] for k in ("key", "system", "system_kind", "kind", "pn", "name", "uid", "state")},
                           "before": a["lines"], **_after(a["lines"], rs)})
    seen, uniq = set(), []
    for a in actions:
        k = (a["kind"], a["system"], a.get("assembly"), a["pn"] if a["kind"] != "replace_line" else a["old_pn"],
             a.get("new_pn") or a.get("to_pn"))
        if k not in seen:
            seen.add(k)
            uniq.append(a)
    systems: dict[str, dict] = {}
    for a in uniq:
        s = systems.setdefault(a["system"], {"name": a["system"], "kind": a.get("system_kind") or "", "actions": []})
        s["actions"].append(a)
    for s in systems.values():
        s["actions"].sort(key=lambda a: (KIND_ORDER[a["kind"]], a.get("assembly_pn") or "", a.get("pn") or a.get("old_pn") or ""))
        s["counts"] = {k: sum(1 for a in s["actions"] if a["kind"] == k) for k in KIND_ORDER}
    cur = main_currency([a.get("currency") for a in uniq if a["kind"] == "move_stock" and a.get("currency")],
                        data.get("currency")) or data.get("currency")
    moved = [a for a in uniq if a["kind"] == "move_stock"]
    return {"groups": groups_out,
            "systems": sorted(systems.values(), key=lambda s: (s["kind"] == "ERP", s["name"])),
            "assemblies": sorted(assemblies, key=lambda a: (a["kind"] == "erp", a["system"], a["pn"] or "")),
            "totals": {"groups": len(groups_out), "old_pns": sum(len(g["old_pns"]) for g in groups_out),
                       "bom_lines": sum(1 for a in uniq if a["kind"] == "replace_line"),
                       "assemblies": len(assemblies), "supersede": sum(1 for a in uniq if a["kind"] == "supersede"),
                       "stock_items": len(moved), "currency": cur,
                       "stock_value": round(sum(a.get("value") or 0 for a in moved if (a.get("currency") or cur) == cur), 2),
                       "systems": len(systems)}}


def _after(lines: list[dict], rs: list[dict]) -> dict:
    """The BOM after the replacements; a line whose new part number is already in the BOM is merged into it."""
    swap = {_key(x["old"]): x for x in rs}
    out: dict[str, dict] = {}
    marks = {}
    for ln in lines:
        x = swap.get(_key(ln["pn"]))
        pn, name = (x["new"], x["new_name"]) if x else (ln["pn"], ln["name"])
        k = _key(pn)
        if k in out:
            out[k]["qty"] += ln["qty"]
            marks[k] = "merged"
        else:
            out[k] = {"pn": pn, "name": name, "qty": ln["qty"], "was": ln["pn"] if x else None}
            marks[k] = "replaced" if x else "same"
    after = [{**v, "change": marks[k]} for k, v in out.items()]
    return {"after": after, "changes": sum(1 for v in after if v["change"] != "same")}


def build_package(res: dict, decisions: dict) -> dict:
    return combine(package_data(res), approved(res, decisions))


# ---------------------------------------------------------------------------------------------------- exports
ACTION_HE = {"replace_line": "החלפת שורה בעץ מוצר", "supersede": "ביטול מק\"ט (מוחלף)", "move_stock": "העברת מלאי וסגירה לרכש"}


def action_text(a: dict) -> str:
    if a["kind"] == "replace_line":
        return f"בעץ המוצר של {a['assembly_pn']}: להחליף את {a['old_pn']} ב-{a['new_pn']}"
    if a["kind"] == "supersede":
        return f"לסמן את {a['pn']} כמבוטל, מוחלף ב-{a['new_pn']}"
    stock = a.get("stock")
    qty = f"{stock:g} {a.get('uom') or ''}".strip() if isinstance(stock, (int, float)) else "את המלאי"
    return f"להעביר {qty} מ-{a['pn']} ל-{a['to_pn']}, ולסגור את {a['pn']} לרכש חדש"


def rows_of(pkg: dict) -> list[dict]:
    """One flat row per action (CSV, Excel)."""
    out = []
    for s in pkg["systems"]:
        for i, a in enumerate(s["actions"], start=1):
            out.append({"system": s["name"], "n": i, "kind": ACTION_HE[a["kind"]], "action": action_text(a),
                        "assembly": a.get("assembly_pn") or "", "assembly_name": a.get("assembly_name") or "",
                        "old_pn": a.get("old_pn") or a.get("pn") or "", "new_pn": a.get("new_pn") or a.get("to_pn") or "",
                        "qty": a.get("qty") if a["kind"] == "replace_line" else a.get("stock"),
                        "value": a.get("value") if a["kind"] == "move_stock" else None,
                        "group": a.get("group_id") or ""})
    return out


CSV_COLS = [("system", "מערכת"), ("n", "#"), ("kind", "סוג"), ("action", "פעולה"), ("assembly", "מכלול"),
            ("assembly_name", "שם המכלול"), ("old_pn", "מק\"ט ישן"), ("new_pn", "מק\"ט שנשאר"), ("qty", "כמות"),
            ("value", "שווי מלאי"), ("group", "קבוצת כפילות")]


def _cell(v) -> str:
    """CSV text; a value that a spreadsheet would read as a formula is kept as text."""
    if v is None:
        return ""
    s = f"{v:g}" if isinstance(v, float) else str(v)
    return "'" + s if s[:1] in ("=", "+", "-", "@", "\t", "\r") else s


def package_csv(pkg: dict) -> str:
    import csv
    import io
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([h for _, h in CSV_COLS])
    for r in rows_of(pkg):
        w.writerow([_cell(r[k]) for k, _ in CSV_COLS])
    return "﻿" + buf.getvalue()       # with a byte order mark, Excel opens Hebrew text correctly


def write_package_xlsx(pkg: dict, path, *, run: str = "", created: str = "", user: str = "", demo: bool = False):
    import datetime as dt
    from pathlib import Path

    from openpyxl import Workbook
    from openpyxl.styles import Font

    from .excel_he import RIGHT, SUB_FONT, TITLE_FONT, _put, _sheet
    t = pkg["totals"]
    wb = Workbook()
    ws = wb.active
    ws.title = "סיכום"
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    for c in "CDEF":
        ws.column_dimensions[c].width = 16
    ws.cell(1, 1, "חבילת שינוי – איחוד מק\"טים כפולים").font = TITLE_FONT
    ws.cell(2, 1, f"הופק {dt.datetime.now():%d/%m/%Y %H:%M}" + (f" על ידי {user}" if user else "") + (f" · סריקה {run} מ-{created[:16].replace('T', ' ')}" if run else "")).font = Font(name="Arial", size=9, color="6B7280")
    ws.cell(3, 1, "PartAudit לא כותב לשום מערכת. זו רשימת עבודה: את השינויים מבצעים במערכות עצמן, לפי הנהלים שלכם." + (" (נתוני הדגמה)" if demo else "")).font = Font(name="Arial", size=9, italic=True, color="6B7280")
    r = 5
    for label, v in (("קבוצות כפילות שאושרו", t["groups"]), ("מק\"טים שיוחלפו", t["old_pns"]), ("שורות בעצי מוצר להחלפה", t["bom_lines"]),
                     ("מכלולים שעץ המוצר שלהם משתנה", t["assemblies"]), ("רשומות לסימון כמבוטלות", t["supersede"]),
                     ("פריטי ERP להעברת מלאי", t["stock_items"]), (f"שווי המלאי שעובר ({t['currency'] or ''})", t["stock_value"])):
        _put(ws, r, 1, label).font = SUB_FONT
        _put(ws, r, 2, v)
        ws.cell(r, 1).alignment = ws.cell(r, 2).alignment = RIGHT
        r += 1
    r += 1
    for c, h in enumerate(("מערכת", "שורות בעץ מוצר", "ביטול מק\"טים", "העברות מלאי"), start=1):
        ws.cell(r, c, h).font = SUB_FONT
    for s in pkg["systems"]:
        r += 1
        for c, v in enumerate((s["name"], s["counts"]["replace_line"], s["counts"]["supersede"], s["counts"]["move_stock"]), start=1):
            _put(ws, r, c, v).alignment = RIGHT
    for s in pkg["systems"]:
        rows = [[i, ACTION_HE[a["kind"]], action_text(a), a.get("assembly_pn") or "", a.get("assembly_name") or "",
                 a.get("old_pn") or a.get("pn") or "", a.get("new_pn") or a.get("to_pn") or "",
                 a.get("qty") if a["kind"] == "replace_line" else a.get("stock"),
                 a.get("value") if a["kind"] == "move_stock" else None, "", ""] for i, a in enumerate(s["actions"], start=1)]
        _sheet(wb, s["name"][:31], [("#", 5), ("סוג", 20), ("פעולה", 52), ("מכלול", 14), ("שם המכלול", 30), ("מק\"ט ישן", 14),
                                    ("מק\"ט שנשאר", 14), ("כמות", 9), ("שווי מלאי", 12), ("בוצע", 8), ("הערות", 24)], rows,
               note=f"{s['name']}: {len(rows)} פעולות. העמודות 'בוצע' ו'הערות' למעקב.")
    rows = []
    for a in pkg["assemblies"]:
        rows.append([f"{a['pn']} – {a['name']}", a["system"], "", "", None, "", None, ""])
        for b in a["before"]:
            k = pn_key(b["pn"])
            x = next((y for y in a["after"] if (y.get("was") and pn_key(y["was"]) == k) or pn_key(y["pn"]) == k), None)
            change = {"replaced": "מוחלף", "merged": "מאוחד", "same": ""}[x["change"]] if x else ""
            rows.append(["", "", b["pn"], b["name"], b["qty"], x["pn"] if x else "", x["qty"] if x else None, change])
    _sheet(wb, "לפני ואחרי", [("מכלול", 34), ("מערכת", 24), ("רכיב לפני", 14), ("שם", 30), ("כמות לפני", 10), ("רכיב אחרי", 14),
                             ("כמות אחרי", 10), ("שינוי", 10)], rows, note="כל מכלול שעץ המוצר שלו משתנה: הרכיבים לפני השינוי ואחריו.")
    _sheet(wb, "קבוצות", [("קבוצה", 12), ("חלק", 36), ("מפרט", 22), ("המק\"ט שנשאר", 14), ("מק\"טים שיוחלפו", 40)],
           [[g["id"], g["name"], g.get("spec") or "", g["target_pn"], ", ".join(g["old_pns"])] for g in pkg["groups"]])
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
