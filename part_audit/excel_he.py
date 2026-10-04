"""Hebrew, right-to-left Excel workbook of an audit result (with the approvals recorded in the app)."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .decisions import chosen_target, effective_mapping, group_impact, row_decision
from .explain import CATEGORY_HE

HEAD = PatternFill("solid", fgColor="1E3A8A")
HEAD_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
BODY_FONT = Font(name="Arial", size=10)
TITLE_FONT = Font(name="Arial", size=16, bold=True, color="1E3A8A")
SUB_FONT = Font(name="Arial", size=11, bold=True, color="1F2937")
INPUT = PatternFill("solid", fgColor="FFF7D6")
BAND = PatternFill("solid", fgColor="F3F6FC")
THIN = Side(style="thin", color="D5DBE5")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)
RIGHT = Alignment(vertical="top", horizontal="right", readingOrder=2)
DECISION_HE = {"approve": "מאושר", "reject": "נדחה", "review": "לבדיקה", "ignore": "התעלם", "": "ממתין"}
# the words on the buttons the engineer pressed (same as in the app)
ISSUE_DECISION_HE = {"approve": "לתקן", "reject": "לא בעיה", "review": "לבדיקה", "ignore": "להתעלם"}
PAIR_DECISION_HE = {
    "near_identical": {"approve": "לתקנן לחלק אחד", "reject": "חלקים שונים – להשאיר", "review": "לבדיקה"},
    "identical_attr_conflict": {"approve": "טעות רישום – לתקן", "reject": "חלקים שונים באמת", "review": "לבדיקה"},
    "mirror": {"approve": "תקין – ימין ושמאל", "reject": "יש טעות בסימון", "review": "לבדיקה"},
    "conflict": {"approve": "לתקן את המק\"ט", "reject": "תקין", "review": "לבדיקה"},
    "conflict_unverified": {"approve": "לתקן את המק\"ט", "reject": "תקין", "review": "לבדיקה"},
}
SEVERITY_HE = {"high": "גבוהה", "medium": "בינונית", "low": "נמוכה", "info": "מידע"}
STATE_HE = {"released": "משוחרר", "in_work": "בעבודה", "review": "בבדיקה", "obsolete": "מבוטל", "": "—"}


def _put(ws, row: int, col: int, v):
    """Write a cell; text that starts with '=' (names come from other systems) stays text, never a formula."""
    cell = ws.cell(row, col, v)
    if isinstance(v, str) and v.startswith("="):
        cell.data_type = "s"
    return cell


def _target(g: dict, parts: dict, decisions: dict) -> str:
    """The member whose part number stays: the approved choice if there is one, else the proposal."""
    return chosen_target(g, decisions.get(f"group:{g['id']}"), parts)


def _impact(g: dict, parts: dict, target: str) -> dict:
    return group_impact(g, parts[target]["pn"], (parts[target].get("erp") or {}).get("pn"))


def _sheet(wb: Workbook, title: str, headers: list[tuple[str, int]], rows: list[list], first: bool = False,
           note: str = "") -> object:
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title[:31]
    ws.sheet_view.rightToLeft = True
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(name="Arial", size=9, italic=True, color="6B7280")
        ws.cell(1, 1).alignment = RIGHT
        r0 = 2
    for c, (h, w) in enumerate(headers, start=1):
        cell = ws.cell(r0, c, h)
        cell.fill, cell.font, cell.alignment, cell.border = HEAD, HEAD_FONT, WRAP, BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    for i, row in enumerate(rows):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(r0 + 1 + i, c, v)
            if isinstance(v, str) and v.startswith("="):
                cell.data_type = "s"   # names come from other systems: never let one become a formula
            cell.font, cell.alignment, cell.border = BODY_FONT, WRAP, BORDER
            if i % 2 == 1:
                cell.fill = BAND
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    if rows:
        ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def write_workbook(res: dict, decisions: dict, path: str | Path, demo: bool = False, user: str = "") -> Path:
    parts = {p["uid"]: p for p in res["parts"]}
    s = dict(res["summary"])
    cur = s.get("currency", "ILS")
    # money and BOM lines follow the target the engineer chose (where one was chosen)
    impacts = []
    for g in res["groups"]:
        impacts.append((g, _impact(g, parts, _target(g, parts, decisions))))
    s["redundant_stock_value"] = round(sum(i["redundant_stock_value"] for g, i in impacts
                                           if (g.get("currency") or cur) == cur), 2)
    s["bom_lines_to_update"] = sum(i["bom_lines_to_update"] for _, i in impacts)
    wb = Workbook()

    # ---------------------------------------------------------------- summary
    ws = wb.active
    ws.title = "סיכום"
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 60
    ws["A1"] = "דוח בקרת נתוני חלקים"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"סריקה מ-{res['created'].replace('T', ' ')} | הופק {dt.datetime.now():%Y-%m-%d %H:%M}" + \
        (f" | על ידי {user}" if user else "")
    ws["A2"].font = Font(name="Arial", size=10, color="6B7280")
    row = 4
    if demo:
        ws.cell(row, 1, "סביבת הדגמה: הנתונים סינתטיים (מפעל בדוי), והחיבורים הם לסימולטורים מקומיים של המערכות.")
        ws.cell(row, 1).font = Font(name="Arial", size=10, bold=True, color="B45309")
        row += 2
    ws.cell(row, 1, "ממצאים עיקריים").font = SUB_FONT
    row += 1
    kpis = [
        ("חלקים שנסרקו", s["parts"], f"ועוד {s['assemblies']} מכלולים, מ-{s['sources']} מקורות"),
        ("חלקים עם קובץ STEP (נבדקה גיאומטריה)", s["with_step"], "חלקים בלי STEP נבדקו לפי שם, מק\"ט ומפרט בלבד"),
        ("קבוצות כפילות", s["groups"], f"{s['groups_cross_source']} מהן חוצות מערכות; {s['groups_standard']} חלקים תקניים"),
        ("מק\"טים שניתן לבטל (אחרי אישור)", s["redundant_pns"], "מספר המק\"טים בכל קבוצה פחות אחד"),
        (f"שווי מלאי במק\"טים הכפולים ({cur})", s["redundant_stock_value"],
         "כמות במלאי × עלות יחידה מה-ERP, בלי המק\"ט שנשאר בכל קבוצה"),
        ("שורות בעצי מוצר שיושפעו", s["bom_lines_to_update"], "שימושים של מק\"טים שיוחלפו"),
        ("זוגות כמעט זהים (לבדיקה)", s["near_pairs"], "מועמדים לתקינה – לא מאוחדים אוטומטית"),
        ("גיאומטריה זהה, מפרט שונה", s["spec_conflicts"], "למשל 2RS/ZZ או 12.9/A2 – לא לאחד"),
        ("אותו מק\"ט עם גיאומטריה שונה", s["pn_conflicts"], ""),
        ("בעיות איכות נתונים (גבוהה / בינונית / נמוכה)",
         f"{s['quality_high']} / {s['quality_medium']} / {s['quality_low']}", ""),
        ("פערים בין PDM ל-ERP", s["erp_issues"], ""),
    ]
    if s.get("assumed_saving_per_year") is not None:
        kpis.append(("חיסכון שנתי משוער (לפי הנחת המשתמש)", s["assumed_saving_per_year"], "מק\"טים מיותרים × עלות שנתית שהוגדרה בהגדרות"))
    for k, v, note in kpis:
        ws.cell(row, 1, k).font = BODY_FONT
        c = ws.cell(row, 2, v)
        c.font = Font(name="Arial", size=11, bold=True)
        if isinstance(v, float):
            c.number_format = "#,##0"
        ws.cell(row, 3, note).font = Font(name="Arial", size=9, color="6B7280")
        for col in (1, 2, 3):
            ws.cell(row, col).alignment = RIGHT
        row += 1
    row += 1
    ws.cell(row, 1, "מקורות נתונים").font = SUB_FONT
    row += 1
    for src in res.get("sources", []):
        st = src.get("stats", {})
        _put(ws, row, 1, src["name"]).font = BODY_FONT
        ws.cell(row, 2, st.get("records") or st.get("erp_items") or 0)
        _put(ws, row, 3, ("סימולטור מקומי; " if src.get("simulated") else "") +
             ("שגיאות: " + "; ".join(src.get("errors") or []) if src.get("errors") else "נקרא בהצלחה"))
        for col in (1, 2, 3):
            ws.cell(row, col).alignment = RIGHT
        row += 1
    row += 1
    ws.cell(row, 1, "איך להשתמש בקובץ").font = SUB_FONT
    row += 1
    for line in ["1. בגיליון 'מיפוי מק\"טים' – לכל שורה בחרו מאושר / נדחה / לבדיקה (תאים צהובים).",
                 "2. 'קבוצות כפילות' מראה לכל קבוצה את כל הרשומות, המערכת שממנה הגיעו, מלאי ושימושים.",
                 "3. 'כמעט זהים' – חלקים דומים מאוד: החלטה הנדסית אם לתקנן לחלק אחד.",
                 "4. התוכנה לא כותבת ל-ERP או ל-PDM. ביצוע השינויים נעשה על ידי הצוות לפי הקובץ המאושר."]:
        ws.cell(row, 1, line).font = BODY_FONT
        ws.cell(row, 1).alignment = RIGHT
        row += 1

    # ---------------------------------------------------------------- mapping
    def dec(kind, tid):
        d = decisions.get(f"{kind}:{tid}") or {}
        return DECISION_HE.get(d.get("decision", ""), d.get("decision", "")), d.get("user", ""), \
            (d.get("created") or "")[:16].replace("T", " "), d.get("note", "")
    rows = []
    for m in effective_mapping(res, decisions):
        d = row_decision(m, decisions)          # the row's own decision or its group's, whichever was made last
        rows.append([m["group_id"], m["old_pn"], m["old_name"], m["old_source"], m["new_pn"], m["new_name"],
                     m["reason_he"], m["why_target_he"], m["confidence"], m["where_used"], m["erp_stock"],
                     DECISION_HE.get(d["decision"], d["decision"]) or "ממתין", d["decision_by"],
                     (d["decision_at"] or "")[:16].replace("T", " "), d["decision_note"]])
    ws = _sheet(wb, "מיפוי מק\"טים", [("קבוצה", 11), ("מק\"ט קיים", 15), ("שם קיים", 32), ("מערכת", 22),
                                      ("מק\"ט יעד", 15), ("שם יעד", 32), ("סיבה", 50), ("למה המק\"ט הזה", 34),
                                      ("ביטחון", 9), ("שימושים", 9), ("מלאי", 9), ("החלטה", 11), ("מאשר", 12),
                                      ("תאריך", 15), ("הערה", 30)], rows,
                note="הצעות בלבד – שום דבר לא נכתב למערכות. רק שורות 'מאושר' הן לביצוע. ההחלטות מהאפליקציה מולאו; ניתן לעדכן בתאים הצהובים.")
    dv = DataValidation(type="list", formula1='"ממתין,מאושר,נדחה,לבדיקה"', allow_blank=True)
    ws.add_data_validation(dv)
    for i in range(len(rows)):
        for col in (12, 13, 14, 15):
            ws.cell(3 + i, col).fill = INPUT
        dv.add(ws.cell(3 + i, 12))
    if rows:
        rng = f"L3:L{2 + len(rows)}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"מאושר"'],
                                                      fill=PatternFill("solid", fgColor="D1FAE5")))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"נדחה"'],
                                                      fill=PatternFill("solid", fgColor="FEE2E2")))

    # ---------------------------------------------------------------- groups
    rows = []
    for g in res["groups"]:
        dg = dec("group", g["id"])
        target = _target(g, parts, decisions)
        erp_by_pn = {e["pn"]: e for e in g["erp"]}
        impact = _impact(g, parts, target)
        for u in g["members"]:
            p = parts[u]
            # the ERP item the part is linked to (the ERP may write the number with other dashes)
            e = erp_by_pn.get((p.get("erp") or {}).get("pn") or p["pn"]) or {}
            rows.append([g["id"], "חלק תקני" if g["category"] == "standard" else "גיאומטריה זהה", g["spec"],
                         ("כן" if u == target else "") + (" (אושר)" if u == target and target != g["master_uid"] else ""),
                         p["pn"], p["name"], p["source"],
                         STATE_HE.get(p["state_class"], p["state"]), p["revision"], p["material"], p["used_in"],
                         e.get("stock"), e.get("unit_cost"), e.get("value"), impact["redundant_stock_value"],
                         g["reason_he"], dg[0] or "ממתין", dg[3]])
    _sheet(wb, "קבוצות כפילות", [("קבוצה", 11), ("סוג", 13), ("מפרט", 16), ("מק\"ט יעד?", 10), ("מק\"ט", 15),
                                 ("שם", 34), ("מערכת", 22), ("מצב", 10), ("רוויזיה", 8), ("חומר", 18),
                                 ("שימושים", 9), ("מלאי", 9), (f"עלות ({cur})", 10), (f"שווי ({cur})", 11),
                                 (f"מלאי כפול בקבוצה ({cur})", 13), ("ראיות", 60), ("החלטה", 11), ("הערה", 28)], rows)

    # ---------------------------------------------------------------- pairs
    def pair_rows(cats):
        out = []
        for p in res["pairs"]:
            if p["category"] not in cats:
                continue
            a, b = parts[p["a"]], parts[p["b"]]
            m = p.get("metrics") or {}
            d = decisions.get(f"pair:{p['id']}") or {}
            label = PAIR_DECISION_HE.get(p["category"], {}).get(d.get("decision", ""), DECISION_HE.get(d.get("decision", ""), ""))
            out.append([CATEGORY_HE.get(p["category"], p["category"]), a["pn"], a["name"], a["source"], b["pn"],
                        b["name"], b["source"], None if m.get("d_volume") is None else round(100 * m["d_volume"], 3),
                        None if m.get("d_area") is None else round(100 * m["d_area"], 3), p["confidence"], p["reason_he"],
                        p["hint_he"], label or "ממתין", d.get("user", ""), d.get("note", "")])
        return out
    hdr = [("סוג", 18), ("מק\"ט א'", 15), ("שם א'", 30), ("מערכת א'", 20), ("מק\"ט ב'", 15), ("שם ב'", 30),
           ("מערכת ב'", 20), ("הפרש נפח %", 10), ("הפרש שטח פנים %", 11), ("ביטחון", 8), ("הסבר", 60),
           ("המלצה", 30), ("החלטה", 16), ("מאשר", 12), ("הערה", 24)]
    _sheet(wb, "כמעט זהים", hdr, pair_rows({"near_identical"}))
    _sheet(wb, "מראה ומפרט שונה", hdr, pair_rows({"mirror", "identical_attr_conflict", "conflict",
                                                    "conflict_unverified"}))

    # ---------------------------------------------------------------- names / quality / ERP
    def issue_dec(iid):
        d = decisions.get(f"issue:{iid}") or {}
        return [ISSUE_DECISION_HE.get(d.get("decision", ""), "") or "ממתין", d.get("user", ""), d.get("note", "")]

    rows = []
    for n in res["names"]:
        ps = [parts[u] for u in n["uids"] if u in parts]
        rows.append([n["type_he"], n.get("pn", ""), " | ".join(n.get("names") or []),
                     " | ".join(f"{p['pn']} ({p['source']})" for p in ps), n.get("message", ""),
                     None if n.get("similarity") is None else round(n["similarity"], 2), *issue_dec(n["id"])])
    _sheet(wb, "שמות ומק\"טים", [("סוג", 26), ("מק\"ט", 15), ("שמות", 60), ("רשומות", 50), ("פירוט", 50),
                                ("דמיון", 8), ("טיפול", 11), ("מי", 12), ("הערה", 24)], rows)
    rows = []
    for q in sorted(res["quality"], key=lambda q: ("high", "medium", "low", "info").index(q["severity"])):
        if q["severity"] == "info":
            continue
        p = parts.get(q["uids"][0]) if q["uids"] else None
        rows.append([q["severity_he"], q["type_he"], p["pn"] if p else "", p["name"] if p else "",
                     p["source"] if p else "", q["message"], *issue_dec(q["id"])])
    _sheet(wb, "איכות נתונים", [("חומרה", 10), ("בדיקה", 30), ("מק\"ט", 15), ("שם", 34), ("מערכת", 22),
                                ("פירוט", 70), ("טיפול", 11), ("מי", 12), ("הערה", 24)], rows)
    rows = []
    for q in sorted(res["erp_issues"], key=lambda q: ("high", "medium", "low", "info").index(q["severity"])):
        p = parts.get(q["uids"][0]) if q["uids"] else None
        rows.append([q["severity_he"], q["type_he"], (p["pn"] if p else q["details"].get("pn", "")),
                     (p["name"] if p else q["details"].get("description", "")), p["source"] if p else "ERP", q["message"],
                     *issue_dec(q["id"])])
    _sheet(wb, "ERP מול PDM", [("חומרה", 10), ("בדיקה", 30), ("מק\"ט", 15), ("שם", 34), ("מערכת", 22),
                               ("פירוט", 70), ("טיפול", 11), ("מי", 12), ("הערה", 24)], rows)

    # ---------------------------------------------------------------- all parts
    rows = []
    for p in res["parts"]:
        e = p.get("erp") or {}
        rows.append([p["pn"], p["name"], "מכלול" if p["kind"] == "assembly" else "חלק", p["source"],
                     STATE_HE.get(p["state_class"], p["state"]), p["revision"], p["material"], p["mass_kg"],
                     "קנייה" if p["make_buy"] == "buy" else ("ייצור" if p["make_buy"] == "make" else ""),
                     "כן" if p["has_step"] else "לא", p["group_id"] or "", p["near_count"], p["used_in"],
                     e.get("pn", ""), e.get("stock"), e.get("unit_cost"), p["native_file"]])
    _sheet(wb, "כל החלקים", [("מק\"ט", 15), ("שם", 36), ("סוג", 8), ("מערכת", 22), ("מצב", 10), ("רוויזיה", 8),
                             ("חומר", 18), ("משקל (ק\"ג)", 10), ("ייצור/קנייה", 10), ("STEP", 6), ("קבוצה", 11),
                             ("כמעט זהים", 9), ("שימושים", 9), ("מק\"ט ERP", 15), ("מלאי", 9), (f"עלות ({cur})", 10),
                             ("קובץ", 40)], rows)
    rows = [[x.get("source", ""), x.get("file", ""), x.get("reason", "")] for x in res.get("skipped", [])]
    for src in res.get("sources", []):
        for e in src.get("errors") or []:
            rows.append([src["name"], "", e])
    _sheet(wb, "שגיאות ודילוגים", [("מקור", 26), ("קובץ", 50), ("סיבה", 60)], rows)
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
