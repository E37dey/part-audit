"""Speed benchmark: run the full audit on N random ABC STEP files (natural mix, incl. multi-body
and broken files) with a cold cache and record wall time, per-file times, errors and peak memory.

    python tools/bench.py --abc RAW/abc0 RAW/abc1 RAW/abc2 --n 5000 --workers 2 --out out/bench
"""
from __future__ import annotations

import argparse
import json
import os
import platform
import random
import resource
import shutil
import subprocess
import sys
import time
from pathlib import Path


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--abc", nargs="+", required=True)
    ap.add_argument("--n", type=int, default=5000)
    ap.add_argument("--workers", type=int, default=os.cpu_count() or 2)
    ap.add_argument("--timeout", type=float, default=60)
    ap.add_argument("--seed", type=int, default=5)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    out = Path(a.out).resolve()
    data = out / "files"
    if out.exists():
        shutil.rmtree(out)
    data.mkdir(parents=True)
    pool = sorted(str(p) for d in a.abc for p in Path(d).rglob("*.step"))
    rng = random.Random(a.seed)
    pick = rng.sample(pool, a.n)
    total_bytes = 0
    for i, src in enumerate(pick):
        sub = data / f"{i // 500:02d}"
        sub.mkdir(exist_ok=True)
        os.symlink(src, sub / Path(src).name)
        total_bytes += os.path.getsize(src)
    cmd = [sys.executable, "-m", "part_audit", "scan", str(data), "--workers", str(a.workers), "--timeout",
           str(a.timeout), "--out", str(out / "bench_report.xlsx"), "--json", str(out / "bench_result.json"),
           "--cache", str(out / "cache"), "-q"]
    t0 = time.perf_counter()
    proc = subprocess.run(cmd, capture_output=True, text=True, cwd=Path(__file__).resolve().parents[1])
    wall = time.perf_counter() - t0
    peak_kb = resource.getrusage(resource.RUSAGE_CHILDREN).ru_maxrss
    if proc.returncode:
        print(proc.stdout[-3000:], proc.stderr[-3000:])
        raise SystemExit("audit failed")
    res = json.loads((out / "bench_result.json").read_text())
    tm = res["timings"]
    from openpyxl import load_workbook
    wb = load_workbook(out / "bench_report.xlsx", read_only=True)
    errs: dict[str, int] = {}
    for row in list(wb["Errors_Skipped"].iter_rows(values_only=True))[1:]:
        errs[row[2]] = errs.get(row[2], 0) + 1
    sizes = sorted(os.path.getsize(p) for p in pick)
    summary = {
        "files": a.n, "total_MB": round(total_bytes / 1e6, 1),
        "size_kb_median": round(sizes[len(sizes) // 2] / 1024, 1), "size_kb_max": round(sizes[-1] / 1024, 1),
        "workers": a.workers, "cpu_count": os.cpu_count(), "timeout_s": a.timeout,
        "wall_s_total": round(wall, 1), "wall_min_total": round(wall / 60, 2),
        "stage_s": {k: round(v, 1) for k, v in tm.items() if k.endswith("_s")},
        "geometry_per_file_s": {k.replace("geometry_per_file_", ""): round(v, 3) for k, v in tm.items()
                                if k.startswith("geometry_per_file_")},
        "seconds_per_file_wall": round(wall / a.n, 3),
        "errors_by_reason": errs, "flagged_pairs": len(res["pairs"]), "duplicate_groups": len(res["groups"]),
        "peak_rss_MB_largest_child": round(peak_kb / 1024, 1),
        "machine": {"platform": platform.platform(), "python": platform.python_version(),
                    "processor": platform.processor() or "n/a"},
        "note": "cold cache; files sampled at random from ABC chunks (<= 3 MB each), natural mix of valid, "
                "multi-body and broken models",
    }
    (out / "bench_summary.json").write_text(json.dumps(summary, indent=1))
    shutil.rmtree(out / "cache", ignore_errors=True)
    print(json.dumps(summary, indent=1))


if __name__ == "__main__":
    main()
