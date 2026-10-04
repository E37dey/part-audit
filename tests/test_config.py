"""Configuration control (products, BOM trees and what is found on each item), the change package built from the
approved decisions, the examples of the "part numbers 101" screen, and their API - on a small hand-made scan."""
import copy
import json
import shutil
import subprocess
from pathlib import Path

import pytest

from part_audit.config import learn_examples, product_list, product_tree
from part_audit.package import approved, build_package, combine, package_csv, package_data, rows_of


def _p(uid, pn, source, kind="part", erp=None, **kw):
    sk = {"PDM A": "swpdm_web", "PLM B": "windchill"}[source]
    return {"uid": uid, "pn": pn, "name": kw.pop("name", f"חלק {pn}"), "kind": kind, "source": source, "source_kind": sk,
            "state": kw.pop("state", "Released"), "state_class": kw.pop("state_class", "released"), "revision": "A",
            "erp": erp, "group_id": kw.pop("group_id", None), "has_step": True, "spec_key": "", "near_count": 0, **kw}


def _erp(pn, stock, cost):
    return {"pn": pn, "description": f"פריט {pn}", "stock": stock, "unit_cost": cost, "currency": "ILS", "status": "פעיל"}


@pytest.fixture
def res():
    parts = [
        _p("a:1", "900-1000", "PDM A", erp=_erp("900-1000", 100, 2.0), group_id="G1"),
        _p("b:1", "0000555", "PLM B", erp=_erp("0000555", 50, 2.2), group_id="G1"),
        _p("a:2", "200-1", "PDM A", state="In Work", state_class="in_work", material=""),
        _p("x:1", "200-2", "PDM A"),
        _p("a:10", "100-1", "PDM A", kind="assembly", erp=_erp("100-1", 0, 0)),
        _p("b:10", "0000999", "PLM B", kind="assembly", erp=_erp("0000999", 0, 0)),
        _p("a:100", "100-9", "PDM A", kind="assembly", name="מכונה"),
    ]
    wu = [{"uid": "a:10", "pn": "100-1", "name": "מכלול 1", "source": "PDM A", "qty": 4.0, "via": "900-1000", "state": "Released"},
          {"uid": "b:10", "pn": "0000999", "name": "מכלול 2", "source": "PLM B", "qty": 2.0, "via": "0000555", "state": "Released"},
          {"uid": "b:10", "pn": "0000999", "name": "מכלול 2", "source": "PLM B", "qty": 1.0, "via": "900-1000", "state": "Released"},
          {"uid": "", "pn": "0000999", "name": "מכלול 2", "source": "ERP", "qty": 3.0, "via": "0000555", "erp_pn": "0000555", "state": ""},
          {"uid": "", "pn": "100-1", "name": "מכלול 1", "source": "ERP", "qty": 4.0, "via": "900-1000", "erp_pn": "900-1000", "state": ""}]
    g = {"id": "G1", "members": ["a:1", "b:1"], "pns": ["0000555", "900-1000"], "master_uid": "a:1", "master_pn": "900-1000",
         "category": "geometry", "reason_he": "אותה גיאומטריה", "confidence": 0.97, "spec": "", "name": "תושבת", "where_used": wu,
         "erp": [{"pn": "900-1000", "stock": 100, "unit_cost": 2.0, "value": 200.0, "currency": "ILS", "is_master": True},
                 {"pn": "0000555", "stock": 50, "unit_cost": 2.2, "value": 110.0, "currency": "ILS", "is_master": False}],
         "redundant_stock_value": 110.0, "bom_lines_to_update": 2, "currency": "ILS"}
    return {"summary": {"parts": 4, "currency": "ILS"}, "parts": parts, "groups": [g],
            "sources": [{"id": "s1", "name": "PDM A", "kind": "swpdm_web"}, {"id": "s3", "name": "Priority", "kind": "priority"}],
            "bom": {"a:100": [{"child": "a:10", "qty": 1.0}, {"child": "b:10", "qty": 2.0}],
                    "a:10": [{"child": "a:1", "qty": 4.0}, {"child": "a:2", "qty": 2.0}],
                    "b:10": [{"child": "b:1", "qty": 2.0}, {"child": "a:1", "qty": 1.0}]},
            "erp_bom": [{"parent": "100-1", "child": "900-1000", "qty": 4.0}, {"parent": "100-1", "child": "200-1", "qty": 2.0},
                        {"parent": "100-1", "child": "300-1", "qty": 1.0}, {"parent": "0000999", "child": "0000555", "qty": 3.0}],
            "pairs": [{"a": "a:1", "b": "b:1", "category": "identical", "reason_he": "זהים", "confidence": 0.97},
                      {"a": "a:2", "b": "x:1", "category": "near_identical", "reason_he": "כמעט", "confidence": 0.8}],
            "mapping": [{"id": "M1", "group_id": "G1", "old_pn": "0000555", "old_uid": "b:1", "new_pn": "900-1000", "new_uid": "a:1",
                         "old_name": "", "old_source": "PLM B", "new_name": "", "reason_he": "", "confidence": 0.97, "where_used": 2,
                         "erp_stock": 50, "status": ""}],
            "quality": [{"type": "missing_material", "type_he": "חסר חומר", "severity": "medium", "uids": ["a:2"], "message": "אין חומר", "id": "Q1"},
                        {"type": "bom_state", "type_he": "רכיב לא משוחרר", "severity": "high", "uids": ["a:2"], "message": "בעבודה", "id": "Q2"},
                        {"type": "no_step", "type_he": "אין STEP", "severity": "info", "uids": ["x:1"], "message": "", "id": "Q3"}],
            "erp_issues": [{"type": "erp_bom_mismatch", "type_he": "עץ מוצר שונה", "severity": "medium", "uids": ["b:10"], "message": "", "id": "E1"}]}


def _approve(gid, master="", n=1, decision="approve"):
    return {f"group:{gid}": {"decision": decision, "master_uid": master, "user": "דנה", "created": "2026-10-04T10:00:00", "id": n}}


# ------------------------------------------------------------------------------------------------ configuration control
def test_products_flags_and_health(res):
    pl = product_list(res, {})
    assert [p["pn"] for p in pl] == ["100-9"]
    h = pl[0]["health"]
    assert h["items"] == 6 and h["lines"] == 7 and h["attention"] == 3 and h["ok"] == 3
    assert h["by_flag"] == {"duplicate": 1, "state": 1, "erp": 1, "quality": 1, "near": 1}
    tree = product_tree(res, "a:100", {})["tree"]
    a1, a2 = tree["children"]
    p1, p3 = a1["children"]
    p2 = a2["children"][0]
    assert [f["type"] for f in p1["flags"]] == ["duplicate"] and p1["flags"][0]["stays"]
    d = p2["flags"][0]
    assert d["type"] == "duplicate" and not d["stays"] and d["severity"] == "high" and d["target_pn"] == "900-1000"
    assert [f["type"] for f in p3["flags"]] == ["state", "quality", "near"]        # an "info" finding is not a flag
    assert [f["type"] for f in a2["flags"]] == ["erp"]
    assert product_tree(res, "a:1", {}) is None                                    # a part is not a product


def test_an_approved_target_changes_the_duplicate_flags(res):
    tree = product_tree(res, "a:100", _approve("G1", master="b:1"))["tree"]
    p1 = tree["children"][0]["children"][0]
    p2 = tree["children"][1]["children"][0]
    assert not p1["flags"][0]["stays"] and p1["flags"][0]["target_pn"] == "0000555" and "אושר" in p1["flags"][0]["text"]
    assert p2["flags"][0]["stays"]


def test_bom_compared_with_the_erp_line_by_line(res):
    tree = product_tree(res, "a:100", {})["tree"]
    a1, a2 = tree["children"]
    st = {r["pn"] or r["erp_pn"]: r["status"] for r in a1["erp_bom"]["rows"]}
    assert st == {"900-1000": "same", "200-1": "same", "300-1": "erp_only"} and a1["erp_bom"]["differences"] == 1
    st2 = {r["pn"]: (r["status"], r["pdm_qty"], r["erp_qty"]) for r in a2["erp_bom"]["rows"]}
    assert st2 == {"0000555": ("qty", 2.0, 3.0), "900-1000": ("pdm_only", 1.0, None)}
    assert tree["erp_bom"]["has_erp_bom"] is False and tree["erp_bom"]["differences"] is None


def test_learn_examples_follow_a_part_with_another_part_number(res):
    ex = learn_examples(res, {})
    assert ex["part_uid"] in ("a:1", "b:1") and ex["group_id"] == "G1" and ex["product_uid"] == "a:100"
    assert ex["counts"]["products"] == 1 and ex["counts"]["approved_groups"] == 0


# ------------------------------------------------------------------------------------------------ change package
def test_package_from_an_approved_group(res):
    pkg = build_package(res, _approve("G1"))
    assert pkg["totals"] == {"groups": 1, "old_pns": 1, "bom_lines": 2, "assemblies": 2, "supersede": 1, "stock_items": 1,
                             "currency": "ILS", "stock_value": 110.0, "systems": 2}
    by = {s["name"]: s for s in pkg["systems"]}
    assert list(by) == ["PLM B", "Priority"]                                       # the ERP last
    plm = [(a["kind"], a.get("assembly_pn"), a.get("old_pn") or a.get("pn"), a.get("new_pn")) for a in by["PLM B"]["actions"]]
    assert plm == [("replace_line", "0000999", "0000555", "900-1000"), ("supersede", None, "0000555", "900-1000")]
    erp = by["Priority"]["actions"]
    assert [(a["kind"], a.get("qty") or a.get("stock")) for a in erp] == [("replace_line", 3.0), ("move_stock", 50)]
    asm = {a["key"]: a for a in pkg["assemblies"]}
    assert asm["pdm:b:10"]["after"] == [{"pn": "900-1000", "name": "חלק 900-1000", "qty": 3.0, "was": "0000555", "change": "merged"}]
    erp_asm = next(a for a in pkg["assemblies"] if a["kind"] == "erp")
    assert erp_asm["pn"] == "0000999" and erp_asm["after"][0]["pn"] == "900-1000" and erp_asm["changes"] == 1
    assert build_package(res, {})["totals"]["groups"] == 0
    assert build_package(res, _approve("G1", decision="review"))["totals"]["groups"] == 0


def test_package_follows_the_chosen_target(res):
    pkg = build_package(res, _approve("G1", master="b:1"))
    t = pkg["totals"]
    assert t["bom_lines"] == 3 and t["stock_value"] == 200.0 and pkg["groups"][0]["target_pn"] == "0000555"
    asm = {a["key"]: a for a in pkg["assemblies"]}
    assert [(x["pn"], x["qty"], x["change"]) for x in asm["pdm:a:10"]["after"]] == [("0000555", 4.0, "replaced"), ("200-1", 2.0, "same")]
    assert [(x["pn"], x["qty"], x["change"]) for x in asm["pdm:b:10"]["after"]] == [("0000555", 3.0, "merged")]
    assert sorted((a["kind"], a["pn"]) for a in pkg["assemblies"]) == [("erp", "100-1"), ("pdm", "0000999"), ("pdm", "100-1")]


def test_a_row_can_be_approved_alone_and_a_later_group_rejection_wins(res):
    row = {"mapping:M1": {"decision": "approve", "user": "דנה", "created": "2026-10-04T10:00:00", "id": 1}}
    assert approved(res, row) == [("G1", "a:1", ["0000555"])]
    later = {**row, **_approve("G1", n=2, decision="reject")}
    assert approved(res, later) == []


def test_csv_keeps_formulas_out_and_rows_are_complete(res):
    r2 = copy.deepcopy(res)
    r2["parts"][1]["pn"] = "=HYPERLINK(1)"
    r2["groups"][0]["pns"] = ["=HYPERLINK(1)", "900-1000"]
    r2["mapping"][0]["old_pn"] = "=HYPERLINK(1)"
    for w in r2["groups"][0]["where_used"]:
        if w["via"] == "0000555":
            w["via"] = "=HYPERLINK(1)"
    pkg = build_package(r2, _approve("G1"))
    text = package_csv(pkg)
    assert text.startswith("﻿") and "'=HYPERLINK(1)" in text and ",=HYPERLINK" not in text
    assert len(rows_of(pkg)) == sum(len(s["actions"]) for s in pkg["systems"])


@pytest.mark.skipif(shutil.which("node") is None, reason="Node.js is not installed")
def test_the_online_copy_puts_the_same_package_together(res, tmp_path):
    """api.js combines the exported package data in the browser; it must give what package.py gives."""
    api = Path(__file__).resolve().parents[1] / "part_audit" / "ui" / "js" / "api.js"
    data = package_data(res)
    cases = [approved(res, _approve("G1")), approved(res, _approve("G1", master="b:1"))]
    (tmp_path / "in.json").write_text(json.dumps({"data": data, "cases": cases}, ensure_ascii=False), encoding="utf-8")
    script = tmp_path / "run.mjs"
    script.write_text(f"""globalThis.window = {{}}; globalThis.location = {{ origin: "http://x" }};
const m = await import({json.dumps(api.as_uri())});
const fs = await import("node:fs");
const x = JSON.parse(fs.readFileSync({json.dumps(str(tmp_path / "in.json"))}, "utf8"));
console.log(JSON.stringify(x.cases.map((c) => m._combinePackage(x.data, c))));
""", encoding="utf-8")
    out = subprocess.run(["node", str(script)], capture_output=True, text=True, check=True).stdout
    got = json.loads(out)
    want = [combine(data, c) for c in cases]
    assert got == json.loads(json.dumps(want, ensure_ascii=False))


# ------------------------------------------------------------------------------------------------ the API
@pytest.fixture
def client(res, tmp_path):
    from fastapi.testclient import TestClient

    from part_audit.server.app import create_app
    from part_audit.server.store import Store
    st = Store(tmp_path / "ws")
    rid = st.new_run()
    (st.run_dir(rid) / "result.json").write_text(json.dumps(res, ensure_ascii=False), encoding="utf-8")
    st.finish_run(rid, "done", res["summary"])
    with TestClient(create_app(tmp_path / "ws", factory_dir=None, demo=False), base_url="http://127.0.0.1:8765") as c:
        yield c


def test_api_products_learn_and_package(client):
    pl = client.get("/api/runs/latest/products").json()["products"]
    assert [p["uid"] for p in pl] == ["a:100"]
    t = client.get("/api/runs/latest/products/a:100").json()
    assert t["health"]["attention"] == 3 and t["flag_he"]["erp"]
    assert client.get("/api/runs/latest/products/nope").status_code == 404
    assert client.get("/api/learn").json()["group_id"] == "G1"
    assert client.get("/api/runs/latest/package").json()["totals"]["groups"] == 0
    assert client.post("/api/decisions", json={"target_type": "group", "target_id": "G1", "decision": "approve"}).json()["ok"]
    pkg = client.get("/api/runs/latest/package").json()
    assert pkg["totals"]["bom_lines"] == 2 and pkg["run_id"]
    csv_ = client.get("/api/runs/latest/package.csv")
    assert csv_.status_code == 200 and "0000555" in csv_.text and csv_.headers["content-type"].startswith("text/csv")
    x = client.get("/api/runs/latest/package.xlsx")
    assert x.status_code == 200
    from io import BytesIO

    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(x.content))
    assert wb.sheetnames == ["סיכום", "PLM B", "Priority", "לפני ואחרי", "קבוצות"]
    assert wb["PLM B"].sheet_view.rightToLeft and wb["PLM B"]["C3"].value.startswith("בעץ המוצר של 0000999")
    assert client.get("/api/runs/latest/package.pdf").status_code == 404
