"""End to end on a small plant: a folder of STEP files + an ERP export, through the pipeline, the decisions, the
Excel export and the local web API (no simulators involved)."""
import json
import random
import shutil
import time

import pytest

pytest.importorskip("OCP")
pytest.importorskip("fastapi")

import cadops  # noqa: E402  (tools/cadops.py)
from part_audit.decisions import effective_mapping  # noqa: E402
from part_audit.excel_he import write_workbook  # noqa: E402
from part_audit.project import ProjectSettings, audit_snapshot, collect  # noqa: E402

ERP_CSV = """מק"ט,תיאור,מלאי,עלות,סטטוס,חומר
100-0001,תושבת פלדה,10,50,פעיל,פלדה
100-0002,BRACKET STEEL,5,48,פעיל,STEEL
100-0003,תושבת עם חור,2,55,פעיל,פלדה
200-0001,פלטה אלומיניום,3,20,פעיל,אלומיניום
999-9999,פריט ללא מודל,7,5,פעיל,
"""


@pytest.fixture(scope="module")
def plant(tmp_path_factory):
    root = tmp_path_factory.mktemp("plant")
    cad = root / "cad"
    (cad / "old").mkdir(parents=True)
    base = cadops.chiral_test_part()
    rng = random.Random(4)
    cadops.write_step(base, str(cad / "100-0001 - תושבת.step"))
    cadops.write_step(cadops.transform(base, cadops.random_rigid(rng)), str(cad / "100-0002 - BRACKET.step"))
    hole, _ = cadops.drill_hole(base, random.Random(2), 0.08)
    cadops.write_step(hole, str(cad / "100-0003 - תושבת עם חור.step"))
    cadops.write_step(cadops.plate_with_hole(), str(cad / "200-0001 - פלטה.step"))
    shutil.copy(cad / "100-0001 - תושבת.step", cad / "old" / "100-0001_revA.stp")
    (cad / "old" / "מסגרת.dwg").write_bytes(b"AC1027")
    (cad / "broken.step").write_text("ISO-10303-21;\nHEADER;\nENDSEC;\nDATA;\nENDSEC;\nEND-ISO-10303-21;\n")
    shutil.copy(cad / "200-0001 - פלטה.step", cad / "=HYPERLINK(1).step")     # a name that looks like a formula
    (root / "erp.csv").write_text(ERP_CSV, encoding="utf-8")
    return root


def _sources(plant):
    return [{"id": "cad", "kind": "folder", "name": "תיקיית CAD", "config": {"path": str(plant / "cad")}, "enabled": True},
            {"id": "erp", "kind": "erp_file", "name": "ERP", "config": {"path": str(plant / "erp.csv")}, "enabled": True}]


@pytest.fixture(scope="module")
def result(plant, tmp_path_factory):
    cache = tmp_path_factory.mktemp("cache")
    snap = collect(_sources(plant), cache)
    return audit_snapshot(snap, ProjectSettings(workers=1), cache)


def _uid(res, pn, contains=""):
    return next(p["uid"] for p in res["parts"] if p["pn"] == pn and contains in p["uid"])


def test_pipeline_finds_the_duplicate_and_the_near_part(result):
    s = result["summary"]
    assert s["parts"] == 7 and s["erp_items"] == 5
    by_pns = {tuple(g["pns"]): g for g in result["groups"]}
    g = by_pns[("100-0001", "100-0002")]
    assert g["category"] == "geometry"
    assert g["master_pn"] == "100-0001"                           # more stock, same state -> stays
    assert len(g["members"]) == 3                                  # the old copy with the same PN joins the group
    assert g["redundant_pns"] == 1 and g["redundant_stock_value"] == pytest.approx(5 * 48)
    near = [p for p in result["pairs"] if p["category"] == "near_identical"]
    pns = {(result["parts"][[q["uid"] for q in result["parts"]].index(p["a"])]["pn"],
            result["parts"][[q["uid"] for q in result["parts"]].index(p["b"])]["pn"]) for p in near}
    assert any("100-0003" in x for x in pns)
    assert not any("100-0001" in g["pns"] and "200-0001" in g["pns"] for g in result["groups"])


def test_mapping_has_only_real_part_number_changes(result):
    rows = [m for m in result["mapping"] if m["new_pn"].startswith("100-")]
    assert [(m["old_pn"], m["new_pn"]) for m in rows] == [("100-0002", "100-0001")]


def test_erp_that_writes_the_number_without_dashes(plant, tmp_path):
    """ERP item '1000001' is the CAD's '100-0001': its stock is the target's own, not a duplicate's."""
    erp = tmp_path / "erp.csv"
    erp.write_text(ERP_CSV.replace("100-0001,", "1000001,"), encoding="utf-8")
    srcs = [_sources(plant)[0], {"id": "erp", "kind": "erp_file", "name": "ERP", "config": {"path": str(erp)},
                                 "enabled": True}]
    res = audit_snapshot(collect(srcs, tmp_path / "c"), ProjectSettings(workers=1), tmp_path / "c")
    g = next(x for x in res["groups"] if "100-0002" in x["pns"])
    assert g["master_pn"] == "100-0001"
    assert g["redundant_stock_value"] == pytest.approx(5 * 48)
    assert {e["pn"]: e["is_master"] for e in g["erp"]} == {"1000001": True, "100-0002": False}


def test_two_systems_one_erp_item(tmp_path):
    """Site A writes 100-0001, site B writes 1000001 like the ERP; the ERP BOM uses the item once. Its BOM line is
    listed once and needs no change whichever record stays; the part card finds it through the ERP link."""
    from openpyxl import Workbook
    from fastapi.testclient import TestClient
    from part_audit.server.app import create_app
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    base = cadops.chiral_test_part()
    cadops.write_step(base, str(tmp_path / "a" / "100-0001 - תושבת.step"))
    cadops.write_step(base, str(tmp_path / "b" / "1000001 - תושבת.step"))
    wb = Workbook()
    wb.active.title = "items"
    for row in (["מק\"ט", "תיאור", "מלאי", "עלות"], ["1000001", "תושבת פלדה", 10, 50], ["500-0001", "מכלול", 1, 900]):
        wb.active.append(row)
    bom = wb.create_sheet("BOM")
    for row in (["parent", "child", "qty"], ["500-0001", "1000001", 2]):
        bom.append(row)
    wb.save(tmp_path / "erp.xlsx")
    srcs = [{"id": "a", "kind": "folder", "name": "A", "config": {"path": str(tmp_path / "a")}, "enabled": True},
            {"id": "b", "kind": "folder", "name": "B", "config": {"path": str(tmp_path / "b")}, "enabled": True},
            {"id": "erp", "kind": "erp_file", "name": "ERP", "config": {"path": str(tmp_path / "erp.xlsx")}, "enabled": True}]
    res = audit_snapshot(collect(srcs, tmp_path / "c"), ProjectSettings(workers=1), tmp_path / "c")
    g = next(x for x in res["groups"] if set(x["pns"]) == {"100-0001", "1000001"})
    assert [(w["source"], w["pn"]) for w in g["where_used"]] == [("ERP", "500-0001")]
    assert g["bom_lines_to_update"] == 0 and g["redundant_stock_value"] == 0
    from part_audit.decisions import group_impact
    for uid in g["members"]:                                     # either record can stay: the ERP item is the same
        p = next(x for x in res["parts"] if x["uid"] == uid)
        assert group_impact(g, p["pn"], p["erp"]["pn"])["bom_lines_to_update"] == 0
    # the part card of the record that writes the dashes shows the ERP parent through the link
    app = create_app(tmp_path / "ws", factory_dir=None, demo=False)
    with TestClient(app, base_url="http://127.0.0.1:8765") as c:
        for s in srcs:
            c.post("/api/sources", json=s)
        assert _scan(c)["status"] == "done"
        uid = next(p["uid"] for p in c.get("/api/runs/latest/parts", params={"q": "100-0001"}).json()["rows"])
        assert [b["parent"] for b in c.get(f"/api/runs/latest/parts/{uid}").json()["erp_parents"]] == ["500-0001"]


def test_quality_findings(result):
    types = {q["type"] for q in result["quality"] + result["erp_issues"]}
    assert "erp_without_cad" in types                            # 999-9999 has no model
    assert any(x["reason"] for x in result["skipped"])            # the .dwg is listed as skipped
    assert "geometry_error" in types                              # the broken STEP file is reported


def test_engineer_can_choose_another_target(result):
    g = next(x for x in result["groups"] if "100-0002" in x["pns"])
    other = _uid(result, "100-0002")
    dec = {f"group:{g['id']}": {"decision": "approve", "master_uid": other, "user": "דנה", "note": "", "created": ""}}
    rows = [m for m in effective_mapping(result, dec) if m["group_id"] == g["id"]]
    assert [(m["old_pn"], m["new_pn"]) for m in rows] == [("100-0001", "100-0002")]
    base = next(m for m in result["mapping"] if m["group_id"] == g["id"])
    assert rows[0]["target_overridden"] and rows[0]["id"] != base["id"]
    assert rows[0]["why_target_he"] == "היעד שאושר על ידי דנה"


def test_a_target_counts_only_once_approved(result):
    """A target picked together with 'review' or 'reject' is remembered but changes nothing yet."""
    from part_audit.decisions import master_pn
    g = next(x for x in result["groups"] if "100-0002" in x["pns"])
    other = _uid(result, "100-0002")
    idx = {p["uid"]: p for p in result["parts"]}
    for decision in ("review", "reject"):
        d = {"decision": decision, "master_uid": other, "user": "דנה", "note": "", "created": ""}
        rows = [m for m in effective_mapping(result, {f"group:{g['id']}": d}) if m["group_id"] == g["id"]]
        assert [(m["old_pn"], m["new_pn"]) for m in rows] == [("100-0002", "100-0001")]
        assert master_pn(idx, g, d) == "100-0001"
    assert master_pn(idx, g, {"decision": "approve", "master_uid": other}) == "100-0002"


def test_later_decision_wins_between_row_and_group(result):
    from part_audit.decisions import row_decision
    m = next(x for x in result["mapping"] if x["old_pn"] == "100-0002")
    row_ok = {"decision": "approve", "user": "a", "created": "2026-10-04T10:00:00", "note": ""}
    grp_no = {"decision": "reject", "user": "b", "created": "2026-10-04T11:00:00", "note": ""}
    d = row_decision(m, {f"mapping:{m['id']}": row_ok, f"group:{m['group_id']}": grp_no})
    assert d["decision"] == "reject" and d["inherited"]               # a later group rejection cancels the row
    d = row_decision(m, {f"mapping:{m['id']}": {**row_ok, "created": "2026-10-04T12:00:00"},
                         f"group:{m['group_id']}": grp_no})
    assert d["decision"] == "approve" and not d["inherited"]          # a later row decision is an exception
    # two clicks in the same second: the order they were saved in decides (the store's ids)
    t = "2026-10-04T10:00:00"
    row_ok, grp_no = {**row_ok, "id": 1, "created": t}, {**grp_no, "id": 2, "created": t}
    assert row_decision(m, {f"mapping:{m['id']}": row_ok, f"group:{m['group_id']}": grp_no})["decision"] == "reject"
    assert row_decision(m, {f"mapping:{m['id']}": {**row_ok, "id": 3},
                            f"group:{m['group_id']}": grp_no})["decision"] == "approve"


def test_excel_is_hebrew_rtl_and_carries_decisions(result, tmp_path):
    from openpyxl import load_workbook
    m = result["mapping"][0]
    dec = {f"mapping:{m['id']}": {"decision": "approve", "user": "דנה", "note": "נבדק", "created": "2026-10-04T10:00:00",
                                  "master_uid": ""}}
    out = write_workbook(result, dec, tmp_path / "r.xlsx", demo=False, user="דנה")
    wb = load_workbook(out)
    assert all(ws.sheet_view.rightToLeft for ws in wb.worksheets)
    ws = wb['מיפוי מק"טים']
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=3)]
    row = next(r for r in rows if r[1] == "100-0002")
    assert row[4] == "100-0001" and row[11] == "מאושר" and row[12] == "דנה"
    # a name that starts with '=' stays text (never a live formula)
    cells = [c for r in wb["כל החלקים"].iter_rows(min_row=2) for c in r if isinstance(c.value, str) and c.value.startswith("=")]
    assert cells and all(c.data_type == "s" for c in cells)
    # the same for the names of the systems on the summary sheet
    odd = {**result, "sources": [{**x, "name": "=HYPERLINK(\"x\")"} for x in result["sources"]]}
    wb = load_workbook(write_workbook(odd, {}, tmp_path / "r2.xlsx"))
    cells = [c for r in wb["סיכום"].iter_rows() for c in r if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(cells) == len(result["sources"]) and all(c.data_type == "s" for c in cells)


# ------------------------------------------------------------------------------------------------ web API

@pytest.fixture(scope="module")
def client(plant, tmp_path_factory):
    from fastapi.testclient import TestClient
    from part_audit.server.app import create_app
    app = create_app(tmp_path_factory.mktemp("ws"), factory_dir=None, demo=False)
    with TestClient(app, base_url="http://127.0.0.1:8765") as c:     # the app answers only on this computer
        yield c


def _scan(client):
    jid = client.post("/api/scan").json()["job_id"]
    for _ in range(600):
        j = client.get(f"/api/jobs/{jid}").json()
        if j["status"] != "running":
            return j
        time.sleep(0.2)
    raise AssertionError("scan did not finish")


def test_api_flow(client, plant):
    st = client.get("/api/status").json()
    assert st["demo_available"] is False and st["latest_run"] is None
    assert client.get("/api/runs/latest/summary").status_code == 404
    for s in _sources(plant):
        sid = client.post("/api/sources", json=s).json()["id"]
        t = client.post(f"/api/sources/{sid}/test").json()
        assert t["ok"], t
    j = _scan(client)
    assert j["status"] == "done", j.get("error")
    summ = client.get("/api/runs/latest/summary").json()
    assert summ["summary"]["groups"] == 2 and summ["decisions"]["pending_groups"] == 2   # brackets; plate + '=' copy
    groups = client.get("/api/runs/latest/groups").json()
    gid = next(g["id"] for g in groups if "100-0002" in g["pns"])
    g = client.get(f"/api/runs/latest/groups/{gid}").json()
    assert len(g["parts"]) == 3 and g["pairs"]

    # 3D: mesh, display alignment and deviation map
    a, b = g["parts"][0]["uid"], g["parts"][1]["uid"]
    mesh = client.get(f"/api/mesh/latest/{a}")
    assert mesh.status_code == 200 and mesh.content[:4] == b"PAM1"
    al = client.get("/api/align/latest", params={"a": a, "b": b}).json()
    assert len(al["matrix"]) == 16 and al["mirror"] is False
    cmp = client.get("/api/compare/latest", params={"a": a, "b": b})
    assert cmp.content[:4] == b"PAC2"

    # decisions: a target that is not in the group is refused; a member is accepted and applied
    bad = client.post("/api/decisions", json={"target_type": "group", "target_id": gid, "decision": "approve",
                                              "master_uid": "nope"})
    assert bad.status_code == 400
    target = next(p["uid"] for p in g["parts"] if p["pn"] == "100-0002")
    ok = client.post("/api/decisions", json={"target_type": "group", "target_id": gid, "decision": "approve",
                                             "master_uid": target, "note": "בדיקה"})
    assert ok.json()["saved"] == 1
    g2 = client.get(f"/api/runs/latest/groups/{gid}").json()
    assert g2["decision"] == "approve" and g2["final_master_pn"] == "100-0002" and len(g2["history"]) == 1
    rows = [r for r in client.get("/api/runs/latest/mapping").json() if r["group_id"] == gid]
    assert [(r["old_pn"], r["new_pn"], r["decision"], r.get("inherited")) for r in rows] == [("100-0001", "100-0002", "approve", True)]

    # findings, parts search and a part card
    for kind in ("names", "quality", "erp_issues"):
        assert isinstance(client.get(f"/api/runs/latest/findings/{kind}").json(), list)
    found = client.get("/api/runs/latest/parts", params={"q": "100-000"}).json()
    assert found["total"] == 4                                   # 100-0001 (two files), 100-0002, 100-0003
    card = client.get(f"/api/runs/latest/parts/{a}").json()
    assert card["uid"] == a and "issues" in card and "pairs" in card

    # export
    x = client.get(f"/api/runs/{st['latest_run'] or 'latest'}/export.xlsx")
    assert x.status_code == 200 and x.content[:2] == b"PK"

    # settings round trip, unknown keys ignored
    client.put("/api/settings", json={"user_name": "דנה", "evil": 1})
    s = client.get("/api/settings").json()
    assert s["user_name"] == "דנה" and "evil" not in s

    # the UI is served
    assert "PartAudit" in client.get("/").text


def test_only_this_computer_and_this_page(client):
    assert client.get("/api/status", headers={"host": "attacker.example:8765"}).status_code == 403
    r = client.post("/api/sources", json={"kind": "folder", "name": "x", "config": {"path": "/tmp"}},
                    headers={"origin": "https://evil.example"})
    assert r.status_code == 403
    assert client.post("/api/scan", content=b"x", headers={"content-type": "text/plain"}).status_code == 415
    assert client.get("/api/docs").status_code == 404            # no CDN-backed API docs page
    assert client.get("/api/status", headers={"host": "testserver"}).status_code == 403   # a name DNS could resolve
    # a change must come from the app's own page: not another local port, not a page that says it is cross-site
    body = {"target_type": "issue", "target_id": "x", "decision": "review"}
    for headers in ({"origin": "http://localhost:9999"}, {"origin": "http://127.0.0.1:8766"}, {"origin": "null"},
                    {"host": "127.0.0.1:9999"}, {"sec-fetch-site": "cross-site"}, {"sec-fetch-site": "same-site"}):
        assert client.post("/api/decisions", json=body, headers=headers).status_code == 403, headers
    ok = client.post("/api/decisions", json=body, headers={"origin": "http://127.0.0.1:8765", "sec-fetch-site": "same-origin"})
    assert ok.status_code == 200
    assert client.get("/api/status", headers={"host": "localhost:9999"}).status_code == 200   # reading is not a change


def test_group_decisions_keep_their_target(client):
    groups = client.get("/api/runs/latest/groups").json()
    gid = next(g["id"] for g in groups if "100-0002" in g["pns"])
    g = client.get(f"/api/runs/latest/groups/{gid}").json()
    other = next(p["uid"] for p in g["parts"] if p["pn"] == "100-0002")

    def decide(decision, **kw):
        client.post("/api/decisions", json={"target_type": "group", "target_ids": [gid], "decision": decision, **kw})
        return client.get(f"/api/runs/latest/groups/{gid}").json()
    client.post("/api/decisions", json={"target_type": "group", "target_id": gid, "decision": "approve", "master_uid": other})
    g2 = decide("approve")                     # a later bulk approval keeps the target approved before
    assert g2["decision_master"] == other and g2["final_master_pn"] == "100-0002"
    g2 = decide("review")                      # under review: the pick is remembered, the proposal is in effect
    assert g2["decision_master"] == other and g2["final_master_pn"] == "100-0001"
    g2 = decide("approve")                     # a bulk approval approves what the list shows
    assert g2["decision_master"] == g["master_uid"] and g2["final_master_pn"] == "100-0001"
    assert client.post("/api/decisions", json={"target_type": "group", "target_ids": [gid, gid], "decision": "approve",
                                               "master_uid": other}).status_code == 400


def test_stopped_scan_is_not_kept(client):
    before = client.get("/api/status").json()["latest_run"]
    jid = client.post("/api/scan").json()["job_id"]
    client.post(f"/api/jobs/{jid}/cancel")
    for _ in range(600):
        j = client.get(f"/api/jobs/{jid}").json()
        if j["status"] != "running":
            break
        time.sleep(0.2)
    assert j["status"] == "cancelled"
    assert client.get("/api/status").json()["latest_run"] == before


def test_a_job_that_finished_is_not_reported_as_stopped():
    """Pressing stop after the scan already finished keeps it: the job is done, not cancelled."""
    import threading
    from part_audit.server.app import Jobs
    jobs, go = Jobs(), threading.Event()
    late = jobs.start("scan", lambda progress, stop: (go.wait(5), {"run_id": "r1"})[1])
    jobs.jobs[late]["cancel"] = True
    go.set()
    go2 = threading.Event()
    stopped = jobs.start("scan", lambda progress, stop: (go2.wait(5), None if stop() else {"run_id": "r2"})[1])
    jobs.jobs[stopped]["cancel"] = True
    go2.set()
    for _ in range(100):
        if all(jobs.get(j)["status"] != "running" for j in (late, stopped)):
            break
        time.sleep(0.05)
    assert jobs.get(late)["status"] == "done" and jobs.get(late)["result"] == {"run_id": "r1"}
    assert jobs.get(stopped)["status"] == "cancelled"


def test_priority_token_is_a_secret_and_moves_with_the_address(client):
    sid = client.post("/api/sources", json={"kind": "priority", "name": "P", "config": {
        "service_root": "https://erp.example/odata/Priority/tabula.ini/demo", "username": "TOKEN123456",
        "password": "PAT"}}).json()["id"]
    listed = json.dumps(client.get("/api/sources").json())
    assert "TOKEN123456" not in listed
    store = client.app.state.store
    # the form sends the mask back: kept while the address is the same, dropped when it changes
    client.put(f"/api/sources/{sid}", json={"config": {"service_root": "https://erp.example/odata/Priority/tabula.ini/demo",
                                                       "username": "••••••", "password": "••••••"}})
    assert store.get_source(sid)["config"]["username"] == "TOKEN123456"
    client.put(f"/api/sources/{sid}", json={"config": {"service_root": "https://other.example/odata/Priority/x/y",
                                                       "username": "••••••", "password": "••••••"}})
    assert store.get_source(sid)["config"]["username"] == "" and store.get_source(sid)["config"]["password"] == ""
    client.delete(f"/api/sources/{sid}")


def test_passwords_are_not_returned(client):
    sid = client.post("/api/sources", json={"kind": "windchill", "name": "WC", "config": {
        "base_url": "https://plm.example/Windchill/servlet/odata", "username": "u", "password": "secret"}}).json()["id"]
    listed = json.dumps(client.get("/api/sources").json())
    assert "secret" not in listed
    client.delete(f"/api/sources/{sid}")
